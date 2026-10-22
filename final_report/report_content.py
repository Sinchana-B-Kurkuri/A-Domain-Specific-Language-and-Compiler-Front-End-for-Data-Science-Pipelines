"""Evidence-linked prose for the final university report."""
from __future__ import annotations
import ast
from collections import Counter
from pathlib import Path
import re
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from report_document import OUT, ROOT, field

TITLE='A Domain-Specific Language and Compiler Front-End for Data Science Pipelines'
CHAPTERS=['Introduction','Problem Statement','Objectives','Literature Survey or Background Study','System Design','Methodology','Compiler Design Concepts Used','Implementation','Testing and Results','Conclusion','Future Enhancements']

SYMBOL_FIELDS=[
 ('data_loaded','bool','Whether a load declaration has been accepted.'),('data_cleaned','bool','Whether at least one clean declaration has been accepted.'),
 ('splits','dict','Split names mapped to percentages.'),('models','dict','Model names mapped to algorithm identifiers.'),
 ('load_line','int or None','Source line of the accepted load declaration.'),('split_lines','dict','Definition lines of split names.'),
 ('model_lines','dict','Definition lines of model names.'),('first_train','tuple or None','Name and line of the first accepted training declaration.'),
 ('selected_columns','list or None','Current ordered selection; None means no select declaration yet.'),('dropped_columns','dict','Dropped column names mapped to their first drop lines.')]

MODULES=[
 ('lexer.py','tokenize','Recognizes tokens and records source coordinates.'),('parser.py','Parser.parse_pipeline','Checks grammar and builds PipelineNode.'),
 ('ast_nodes.py','PipelineNode and statement classes','Defines AST dataclasses.'),('semantic.py','analyze and SymbolTable','Checks dependencies and records symbolic state.'),
 ('ir.py','IRInstruction and lower_to_ir','Lowers recognized statements into a flat instruction list.'),('optimizer.py','optimize','Applies four ordered IR transformations and records notes.'),
 ('executor.py','execute','Formats symbolic trace entries from opcodes.'),('cli.py','main','Coordinates the stages, flags and expected compiler-error messages.')]

REFERENCES=[
 ('A. V. Aho, M. S. Lam, R. Sethi, and J. D. Ullman, Compilers: Principles, Techniques, and Tools, 2nd ed. Pearson.', 'https://www.pearson.com/en-us/subject-catalog/p/compilers-principles-techniques-and-tools/P200000003472/9780321486813'),
 ('R. Nystrom, Crafting Interpreters, chapter Parsing Expressions.', 'https://craftinginterpreters.com/parsing-expressions.html'),
 ('Python Software Foundation, re - Regular expression operations, Python documentation.', 'https://docs.python.org/3/library/re.html'),
 ('Python Software Foundation, dataclasses - Data Classes, Python documentation.', 'https://docs.python.org/3/library/dataclasses.html'),
 ('pytest developers, pytest documentation.', 'https://docs.pytest.org/en/stable/'),
 ('Streamlit, Streamlit documentation.', 'https://docs.streamlit.io/'),
 ('Apache Software Foundation, Apache Airflow Architecture Overview.', 'https://airflow.apache.org/docs/apache-airflow/stable/core-concepts/overview.html'),
 ('Luigi contributors, Getting Started, Luigi documentation.', 'https://luigi.readthedocs.io/en/stable/'),
 ('Kubeflow contributors, Kubeflow Pipelines Overview.', 'https://www.kubeflow.org/docs/components/pipelines/overview/'),
 ('scikit-learn developers, Pipeline, scikit-learn API documentation.', 'https://scikit-learn.org/stable/modules/generated/sklearn.pipeline.Pipeline'),
 ('Compiler Design Laboratory Project Instruction Manual, supplied course document, section 24.', ''),
 ('S. B. Kurkuri, Phase 1 Problem Definition and System Design, BCSE307P, Fall Semester 2026-27. Supplied project report.', ''),
 ('S. B. Kurkuri, Phase 2 Core Implementation, BCSE307P, Fall Semester 2026-27. Supplied project report.', ''),
 ('S. B. Kurkuri, Data Science Pipeline DSL source repository and accompanying tests. Local working-tree evidence and recorded Git history.', 'https://github.com/Sinchana-B-Kurkuri/A-Domain-Specific-Language-and-Compiler-Front-End-for-Data-Science-Pipelines')]


def p_many(r,*paragraphs):
    for paragraph in paragraphs:r.p(paragraph)


def front(r,data):
    p=r.doc.add_paragraph('FINAL PROJECT REPORT',style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(45)
    p=r.doc.add_paragraph(TITLE,style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(35);p.paragraph_format.space_after=Pt(30)
    r.p('BCSE307P - Compiler Design Laboratory',align=WD_ALIGN_PARAGRAPH.CENTER)
    r.p('Fall Semester 2026-27',align=WD_ALIGN_PARAGRAPH.CENTER)
    r.p('Phase 3 Final Implementation and Testing',align=WD_ALIGN_PARAGRAPH.CENTER)
    r.p('',align=WD_ALIGN_PARAGRAPH.CENTER)
    p=r.p('Prepared by',align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.space_before=Pt(45)
    r.p('SINCHANA B KURKURI',align=WD_ALIGN_PARAGRAPH.CENTER)
    r.p('24BDS0352',align=WD_ALIGN_PARAGRAPH.CENTER)
    r.p('Slot L49+L50',align=WD_ALIGN_PARAGRAPH.CENTER)
    r.pagebreak();r.front_heading('Abstract')
    total=data['counts']['totals'].get('passed',0);examples=len(data['examples'])
    abstract=(f'Data science pipeline descriptions contain dependencies that can be checked before an execution engine is invoked. '
      'This project implements a domain-specific language and compiler front-end for expressing such descriptions as one named pipeline. '
      'The language supports dataset declarations, cleaning, column selection and dropping, percentage-based partitions, model declarations and evaluation requests. '
      'A regular-expression lexer produces tokens with source coordinates. A hand-written recursive-descent parser constructs a dataclass-based abstract syntax tree. '
      'Semantic analysis traverses the tree with a symbol table to enforce declaration, ordering, column and partition constraints. '
      'A lowering stage produces opcode-and-argument instructions. Four optimizer passes remove instructions after the final evaluation, eliminate unused training declarations, merge adjacent drops and collapse adjacent duplicate cleaning or selection instructions. '
      'A symbolic executor formats the optimized instructions as a readable trace; no dataset contents are read and no numerical model training or accuracy calculation is performed. '
      'A command-line interface exposes intermediate results, while a Streamlit dashboard presents the same compiler stages interactively. '
      f'Validation records {total} passing tests and the outcomes of all {examples} supplied example programs. '
      'Optimizer tests compare symbolic evaluation records, test transformation boundaries and check repeated optimization. '
      'Recorded benchmark measurements are examined by stage, time per statement and instruction reduction. '
      'The results demonstrate an inspectable compiler pipeline within a deliberately restricted language, with explicit limits on execution, schema validation and error recovery.')
    r.p(abstract)
    data['abstract_words']=len(abstract.split())
    for title,instruction in [('Table of Contents','TOC \\o "1-2" \\h \\z \\u'),('List of Figures','TOC \\h \\z \\c "Figure"'),('List of Tables','TOC \\h \\z \\c "Table"')]:
        r.pagebreak();r.front_heading(title)
        p=r.doc.add_paragraph();field(p,instruction,'')
    r.pagebreak();r.front_heading('List of Abbreviations')
    for short,long in [('AST','Abstract Syntax Tree'),('CLI','Command-Line Interface'),('CRLF','Carriage Return and Line Feed'),('CSV','Comma-Separated Values'),('DAG','Directed Acyclic Graph'),('DSL','Domain-Specific Language'),('EBNF','Extended Backus-Naur Form'),('EOF','End of File'),('IR','Intermediate Representation'),('LF','Line Feed'),('LL(1)','Left-to-right predictive parsing with one token of lookahead'),('ML','Machine Learning')]:
        r.p(f'{short} - {long}',align=WD_ALIGN_PARAGRAPH.LEFT)


def chapters_1_to_6(r,d):
    counts=d['counts'];programs=d['pipeline']['programs']
    r.chapter_start(1,CHAPTERS[0],first=True)
    r.h('1.1 Background and problem area')
    p_many(r,
      'A pipeline describes an ordered set of operations whose later steps depend on earlier declarations. Even a short workflow introduces relationships between a dataset, selected columns, named partitions and models. A textual description becomes easier to inspect when these relationships are represented as language rules rather than scattered across execution code.',
      'Compiler construction provides a suitable separation of concerns. Lexical analysis recognizes the vocabulary, parsing establishes grammatical structure, and semantic analysis checks contextual relationships. A valid representation can then be lowered and transformed before an interpreter or another back end consumes it [1]. This project applies those stages to a restricted data science pipeline language.')
    r.h('1.2 Motivation')
    p_many(r,
      'An ordering mistake such as training before defining a training partition can be identified without opening a dataset. Similarly, duplicate model names and percentages outside the accepted range are properties of the program description. Detecting these cases before a concrete runtime is invoked makes the source of a failure easier to isolate.',
      'The project therefore concentrates on inspectable language processing. It does not measure time saved by avoiding real training jobs. The executor is deliberately symbolic, making demonstrations independent of dataset availability and machine learning libraries. The resulting traces expose instruction order without implying numerical computation.')
    r.h('1.3 Language and implementation overview')
    p_many(r,
      'A source file contains one `pipeline` declaration. Its body may contain `load`, `clean`, `select`, `drop`, `split`, `train` and `evaluate` statements. The root declaration names the program; it is not an eighth executable body operation. The parser accepts an empty body, and semantics does not impose a final mandatory-load check on such a body.',
      'The compiler package is `ds_pipeline_dsl`. It is driven by a command-line interface and reused by a Streamlit dashboard. Both expose actual compiler results. The dashboard adds presentation, stage timing and partial-result display without introducing new language statements.')
    r.h('1.4 Report organization')
    r.p('Chapters 2-4 establish the problem, objectives and background. Chapters 5-8 explain design, development, compiler concepts and implementation. Chapter 9 evaluates the results, followed by conclusions and future work. Appendices provide the grammar, diagnostics, listings, examples, reproduction commands and phase evolution.')

    r.chapter_start(2,CHAPTERS[1])
    r.h('2.1 Problem definition')
    p_many(r,
      'The problem is to design and implement a small language processor that accepts a pipeline description, rejects malformed or contextually inconsistent programs with source-referenced diagnostics, and translates accepted programs into an inspectable instruction sequence. The processor must keep recognition, structural analysis, state validation and execution presentation separate.',
      'Validation performed only after a concrete workflow begins can discover description errors later than necessary. This project demonstrates earlier detection for the specific invariants encoded in its semantic analyzer. It does not establish that general-purpose tools lack validation, nor does it provide a universal static analyzer for arbitrary Python programs.')
    r.h('2.2 Functional boundaries')
    r.table('Scope of the implemented processor',['Area','Implemented scope','Boundary'],[
      ('Source language','One named pipeline with seven body statement forms.','No loops, expressions, functions or multiple pipeline blocks.'),
      ('Validation','Token, grammar, declaration, ordering, column-state and split constraints.','No CSV schema lookup or general data-type inference.'),
      ('Transformation','AST lowering and four IR optimizer passes.','No target machine-code generation or arbitrary control-flow analysis.'),
      ('Execution','Readable symbolic trace from IR.','No file-content loading, model fitting, prediction or numeric accuracy.'),
      ('Interaction','CLI and local Streamlit dashboard.','No distributed scheduler or visual node-based pipeline editor.')],[1.3,2.7,2.8],intro='Table {ref} defines the implementation boundary against which the results are evaluated.'.replace('Table {ref}','{ref}'))
    r.h('2.3 Inputs outputs and acceptance conditions')
    r.p('Input is source text, normally stored in a `.pipe` file. Successful processing produces tokens, an AST, a final symbol table, original IR, optimized IR, optimizer notes and trace strings. Expected compiler failures produce the original lexical, syntax or semantic diagnostic. The CLI returns zero for successful compilation and one for its handled compiler and file-reading errors. A request to display only one stage changes printed output, not the compiler stages that are executed.')

    r.chapter_start(3,CHAPTERS[2])
    r.h('3.1 Objectives carried forward from Phase 1')
    r.p('The original proposal specifies grammar design, tokenization, parsing, symbol-table state, semantic validation, IR lowering, optimization, symbolic execution and evaluation [12]. Those objectives remain the basis for the final assessment. Table 3.1 separates implemented dependencies and argument checks from broader type and scope checking that the initial wording suggested.')
    phase=d['phases']['Phase1.docx'];inside=False;objectives=[]
    for item in phase:
        text=item.get('text','')
        if text=='5. Objectives':inside=True;continue
        if inside and item.get('style')=='Heading 1':break
        if inside and item.get('style')=='List Paragraph':objectives.append(text)
    rows=[]
    evidence=['parser.py and grammar appendix','lexer.py and test_lexer.py','parser.py and test_parser.py','SymbolTable and test_semantic.py',
              'analyze; scope and allowed-value limits in Section 3.2','lower_to_ir and test_ir.py','optimize and test_optimizer.py','execute and test_executor.py','pytest log and all example CLI runs']
    for i,objective in enumerate(objectives):
        rows.append((str(i+1),objective,'Partially achieved' if i==4 else 'Achieved',evidence[i]))
    r.table('Original objectives and final status',['No.','Phase 1 objective','Status','Evidence'],rows,[.4,3.8,1.25,1.8],font_size=9.5,
            intro='The objective-by-objective assessment is presented in {ref}; the original objective wording is retained.')
    r.h('3.2 Scope decisions and additional outcomes')
    p_many(r,
      'Semantic dependency checking and argument constraints are implemented, including percentage bounds, duplicate names and selected/dropped-column relationships. Multiple pipeline scopes, a general type system and allowed-value lists for strategy or algorithm names were not pursued. Names such as a custom cleaning strategy remain valid identifiers. This limitation makes the semantic-analysis objective partial rather than implying implementation of a full type checker.',
      'The final language adds `drop` beyond the Phase 1 grammar. The optimizer expands to four explicit passes, and the project includes a benchmark and dashboard. These additions extend the demonstration while retaining symbolic execution. ANTLR4 was considered in the proposal but was not adopted; lexical analysis and parsing remain hand-written.')

    r.chapter_start(4,CHAPTERS[3])
    r.h('4.1 Workflow orchestration and library pipelines')
    p_many(r,
      'Apache Airflow represents workflows as DAGs of tasks with dependencies and a scheduler/executor architecture [7]. Luigi supports dependency resolution and management of batch-task workflows [8]. Their task-level execution concerns are broader than the small source language implemented here. These systems should not be described as having no validation; this project instead makes its own statement-level semantic rules explicit.',
      'Kubeflow Pipelines provides a platform for machine learning workflows [9]. Its pipeline tooling is distinct from this project, which neither deploys containers nor executes machine learning workloads. Scikit-learn Pipeline composes transformers and an optional final predictor through estimator interfaces [10]. It provides execution-oriented composition rather than the specific `.pipe` grammar, AST and source diagnostics implemented by this project.')
    r.table('Related approaches and the project boundary',['Approach','Primary representation','Primary concern','Relation to this project'],[
      ('Airflow','DAGs and tasks','Workflow scheduling and execution.','Task orchestration is outside the implemented DSL.'),
      ('Luigi','Tasks with dependencies','Batch-job dependency management.','Related dependency idea at a different abstraction level.'),
      ('Kubeflow Pipelines','ML workflow components','Workflow compilation and execution infrastructure.','No deployment or concrete ML execution is provided here.'),
      ('scikit-learn Pipeline','Transformers and estimators','Composable fit/transform/predict interfaces.','The present project validates a restricted textual description.'),
      ('This project','Tokens, AST, SymbolTable and linear IR','Compiler-stage inspection and explicit static invariants.','Symbolic execution supports a reproducible compiler demonstration.')],[1.2,1.5,2.0,2.5],font_size=9.5)
    r.h('4.2 Compiler foundations')
    p_many(r,
      'Lexical analysis groups characters into tokens using patterns. The Python `re` module supplies the compiled matching operations used by the scanner [3]. A parser then recognizes sequences of token types. Recursive descent is a top-down approach in which procedures follow grammar structure; it is suitable when the next token selects the required alternative without backtracking [2].',
      'An AST stores semantic structure with less punctuation detail than the source. Dataclasses provide generated initialization and representation methods for the AST and IR objects [4]. A symbol table records contextual information during a traversal. In this language it stores pipeline artifacts and definition lines rather than nested variable scopes.',
      'An intermediate representation separates source structure from transformations and execution. The representation here is a linear list of opcode/argument records, not a register-based target language or a conventional three-address expression form. Optimization correctness is stated relative to the symbolic evaluation-record oracle implemented in the tests, rather than equality of every textual trace line.')
    r.h('4.3 Gap addressed')
    r.p('The project addresses the need for a small, inspectable example of compiler integration in a data science vocabulary. It exposes every representation, makes its invariants testable and produces deterministic traces for accepted examples. Its contribution is the combination of a restricted grammar, contextual validation, an explicit IR and checked transformations. It is not presented as a replacement for the execution platforms reviewed above.')

    r.chapter_start(5,CHAPTERS[4])
    r.h('5.1 Architecture and data flow')
    r.figure('architecture.png','Compiler architecture and representation flow',intro='{ref} shows the successful processing path. Each stage consumes a representation produced earlier in the pipeline.')
    r.p('The AST is retained while semantic analysis constructs a separate SymbolTable. Lowering receives both objects, although the current lowerer uses statement fields and does not consult the symbol-table parameter. Optimization returns a new instruction sequence and may append textual notes to an optional list. The executor returns trace strings rather than printing them.')
    r.h('5.2 Failure propagation')
    r.figure('error_flow.png','First-error control flow',intro='The three expected compiler error classes form the failure paths in {ref}. A failure prevents later compiler stages from running.')
    r.p('The CLI catches the specific compiler exceptions and prints their messages. The dashboard adapter also catches unexpected exceptions, stores a traceback separately and marks subsequent stages skipped. Consequently, dashboard error resilience is broader than the CLI exception boundary; it does not change which programs the compiler accepts.')
    r.h('5.3 Module responsibilities')
    r.table('Compiler modules and responsibilities',['Module','Principal API','Responsibility'],MODULES,[1.3,2.0,3.2],font_size=10)
    r.figure('package_structure.png','Compiler package structure',intro='The module organization in {ref} corresponds to the physical package under `src/ds_pipeline_dsl/`.')
    r.h('5.4 Symbol table and IR design')
    r.p('The symbol table models the state accumulated from earlier accepted statements. Definition-line maps make duplicate diagnostics explain both the current and previous declaration. `selected_columns=None` is meaningful: a drop before any selection is permitted after load, whereas a drop after selection must be drawn from the current selection and must leave at least one selected column. The complete field list appears in Table 7.3.')
    r.p('IRInstruction contains `op: str` and `args: dict`. Every recognized statement is lowered to one instruction in source order. Source lines are stored in the argument dictionary so the optimizer can explain removed or merged operations. The pipeline root itself has no executable opcode.')
    r.h('5.5 Interfaces and presentation design')
    r.p('The CLI uses argparse flags for tokens, AST, original IR, optimized IR and trace. With no display flag, all five display sections are printed, while semantic analysis runs between parsing and lowering without its own CLI section. This distinction explains why six processing stages appear in the dashboard but five headings appear in default CLI output.')
    r.table('Dashboard pages and evidence shown',['Page','Presentation role'],[
      ('Overview','Static language label; grouped pipeline declaration and operations; discovered example count; saved test evidence; grammar reference.'),
      ('Compiler Playground','Example selection, source editing, explicit Compile action, statuses, error highlighting, tokens, AST, semantic state, IR, optimization and trace.'),
      ('Architecture','The compiler diagram, last-run colors, stage descriptions, implementation references and inline live counts.'),
      ('Performance','Recorded environment, benchmark table, stage timing, time-per-statement and IR-size charts, and observations.')],[1.6,4.6],intro='{ref} describes the implemented dashboard pages. It is an inspection interface over the compiler, not another language-processing implementation.')
    r.h('5.6 Project directory organization')
    r.code('src/ds_pipeline_dsl/  Compiler modules\ntests/                Automated tests and semantic oracle\nexamples/             Supplied .pipe programs\nbenchmarks/           Synthetic benchmark driver\ndashboard/            Streamlit UI and pure adapters\n.streamlit/           Local dashboard configuration\ndocs/                 Earlier reports and project documentation\nfinal_report/         Final report, build code, assets and evidence')

    r.chapter_start(6,CHAPTERS[5])
    r.h('6.1 Phased development')
    p_many(r,
      'Phase 1 established the problem, grammar, proposed architecture and a working lexer/parser/AST/semantic prototype [12]. Phase 2 reorganized the implementation into a package and introduced IR, optimization, symbolic execution, CLI access and an automated suite documented as nineteen tests [13]. These phase reports describe intermediate states and are not used as the final API specification.',
      'The final phase adds frontend and semantic hardening, the `drop` construct, safe optimizer passes with diagnostics, more tests, benchmarking and documentation. The current working tree also contains the Streamlit dashboard. Its presence is verified from files and tests; it is not attributed to a commit when the recorded history does not contain that dashboard addition.')
    r.h('6.2 Development history')
    history=d['history']
    r.p(f'The local Git history contains {len(history)} recorded commits, including one merge commit. The sequence provides evidence of development themes and milestones, rather than a complete record of individual editing sessions. Table 6.1 summarizes those themes.')
    r.table('Recorded development milestones',['Commit or range','Verified theme'],[
      ('7e8c390 to e39721e','Initial repository, Phase 1 prototype/design and a merge.'),
      ('4c8e14f to e0ed13e','Package refactoring, obsolete-prototype removal and split-range validation.'),
      ('460041f to 90fe669','IR, an initial optimizer, symbolic executor and CLI.'),
      ('ab9fe03 to 7fb8b1d','Phase 2 implementation/documentation and Review 2 presentation.'),
      ('4e3dd75 and ac9564b','Final frontend/semantic improvements and complete drop integration.'),
      ('5d490ed and e7d2cfe','Safe optimization diagnostics and comprehensive compiler/integration tests.'),
      ('45a0532 and 9eca51a','Benchmarking, evaluation and final-phase documentation.')],[1.65,4.55])
    r.h('6.3 Design decisions')
    p_many(r,
      'Hand-written regular-expression scanning and recursive-descent parsing keep the grammar-to-code relationship explicit. The body alternatives begin with distinct keyword tokens, and optional list repetition is selected by commas. A generated parser was unnecessary for this grammar, so the proposed ANTLR4 alternative was not pursued.',
      'One pipeline per file avoids scope resolution across separate programs. Strategy and algorithm names remain free identifiers to avoid coupling syntax to a particular ML library. First-error reporting limits cascading diagnostics, while the symbolic executor keeps the demonstration independent of data and training dependencies. Each choice reduces the implementation boundary in a way that can be explained and tested.')
    r.h('6.4 Optimizer corrections and regression validation')
    p_many(r,
      'The Phase 2 optimizer captured in commit ab9fe03 sets a flag after encountering an OP_EVAL and removes every subsequent instruction. This is unsafe for a legal sequence containing a second model and evaluation. The final optimizer searches for the last evaluation before trimming the tail. Current regression and sensitivity tests make the lost-second-evaluation case explicit.',
      'The earlier duplicate comparison includes all instruction arguments. Since lowering includes the source line in those arguments, identical operations on different lines can compare unequal. The final duplicate pass excludes line metadata from the semantic comparison while retaining it in surviving instructions and notes.',
      'Current tests exercise the corrected behaviors, compare each pass with an evaluation-record oracle and test repeated optimization. The available commit order does not establish whether tests were written before each implementation edit; the methodology is therefore described as regression-driven validation rather than an unsupported claim about test-first chronology.')
    r.h('6.5 Reconciliation of earlier descriptions')
    r.p('The final Token is a dataclass, despite the Phase 2 text referring to a namedtuple. The processor has dependency/state checks rather than general type inference or multiple-pipeline scope checking. An empty pipeline is accepted, so a mandatory single-load claim would be incorrect. Unused and trailing models may appear in the semantic symbol table because that table describes the original AST; optimization subsequently removes their IR instructions where permitted.')


def grammar_text():
    source=(ROOT/'src/ds_pipeline_dsl/parser.py').read_text(encoding='utf-8')
    doc=ast.get_docstring(ast.parse(source),clean=False)
    return ('pipeline_decl'+doc.split('pipeline_decl',1)[1].split('Lexical Notes:',1)[0]).strip()


def chapters_7_and_8(r,d):
    data=d['pipeline'];basic=data['programs']['valid_basic.pipe'];demo=data['programs']['optimizer_demo.pipe']
    r.chapter_start(7,CHAPTERS[6])
    r.h('7.1 Lexical analysis')
    p_many(r,
      'The scanner compiles the ordered TOKEN_SPEC alternatives into one master regular expression with named groups. `finditer` yields matches from left to right. Identifiers that exactly match a reserved word are converted to the corresponding uppercase keyword token. Keyword recognition therefore does not split a longer identifier such as `trainer` into a keyword and suffix.',
      'Table 7.1 preserves the actual regular expressions and their priority order. STRING precedes BAD_STRING, so a closed literal is recognized before the unterminated-string case. PERCENT precedes NUMBER, allowing the percentage suffix to form one token. EOF is appended explicitly and is not a regex alternative.')
    purposes=['Quoted string literal','Unterminated string detector','Numeric percentage with suffix','Bare numeric literal','Opening brace','Closing brace','Assignment separator','Statement terminator','Column-list separator','Identifier or keyword candidate','Whitespace, silently skipped','Line comment, silently skipped','Unexpected-character detector']
    rows=[(str(i),name,pattern,purpose) for i,((name,pattern),purpose) in enumerate(zip(data['token_spec'],purposes),1)]
    rows.append(('Synthetic','EOF','Empty value; appended after scanning','End-of-file marker'))
    r.table('Lexical patterns in scanner priority order',['Order','Name','Regular expression','Role'],rows,[.65,1.2,2.55,2.1],font_size=9.5)
    r.p('The reserved words are '+', '.join('`'+k+'`' for k in data['keywords'])+'. Whitespace includes spaces, tabs, CR and LF. Line tracking counts LF characters; a CRLF sequence therefore advances one line. Comment text is ignored until LF. The identifier regex is ASCII-based, while comments and strings can contain Unicode characters. The digit patterns use Python regular-expression `\\d`, not a separately enforced ASCII digit alphabet.')
    r.p('The scanner initializes line numbering to one and calculates columns as offsets from the current line start, making columns zero-based. The appended EOF token uses the current line and column zero. Unterminated strings and unexpected characters raise LexicalError. These are lexer errors rather than parser guesses about an invalid token.')
    r.h('7.2 Syntax analysis')
    r.code(grammar_text())
    p_many(r,
      'The grammar above is taken from the parser specification and checked against its statement branches. `parse_pipeline` requires PIPELINE, an identifier and an opening brace; it repeatedly parses statements until a closing brace is found. It then requires the remaining token to be EOF, with a specific diagnostic for a second pipeline declaration.',
      'The body alternatives have pairwise-disjoint FIRST sets: LOAD, CLEAN, SELECT, DROP, SPLIT, TRAIN and EVALUATE. The repeated statement sequence stops on RBRACE. A column-list continuation is selected by COMMA and otherwise the parser expects SEMI. These choices require only one lookahead token and no backtracking. There is no left-recursive production. This is the relevant LL(1) justification for the implemented grammar, rather than a claim about arbitrary grammar extensions.',
      'The split-name position accepts IDENT or TRAIN. The word `test` is already an ordinary identifier, so it needs no separate reserved-token alternative. Other reserved words are not accepted where a normal identifier is expected. Percentages are tokenized with the suffix present; a bare NUMBER at a split value leads to an expected-PERCENT syntax diagnostic.',
      'The parser uses `peek`, `advance` and `expect` to manage its token position. `expect` reports the required token type and the type/value found at the current token. A missing semicolon is often detected at the following statement token, so its error line can be later than the line where the omission occurred. This behavior is demonstrated by the supplied syntax-error example in Chapter 9.')
    r.h('7.3 Abstract syntax tree')
    classes=[]
    tree=ast.parse((ROOT/'src/ds_pipeline_dsl/ast_nodes.py').read_text(encoding='utf-8'))
    for node in tree.body:
        if isinstance(node,ast.ClassDef):
            attrs=[f'{n.target.id}: {ast.unparse(n.annotation)}' for n in node.body if isinstance(n,ast.AnnAssign)]
            classes.append((node.name,'; '.join(attrs)))
    r.table('AST node classes and declared fields',['Node class','Fields'],classes,[1.6,4.6],font_size=10)
    r.figure('ast_basic.png','AST produced for the basic valid pipeline',intro='{ref} shows the six statement nodes returned for `valid_basic.pipe`, retaining their actual fields and line numbers.',width=5.6)
    r.p('The root stores the pipeline name and ordered statement list. Statement nodes keep the fields needed by later stages and the source line at which the keyword was encountered. The AST omits braces, commas and semicolons after they have served their parsing role. It therefore represents structure rather than a complete concrete syntax tree. Dataclass annotations document field intent but do not themselves implement runtime type enforcement.')
    r.h('7.4 Symbol table and semantic analysis')
    r.table('SymbolTable fields',['Field','Stored form','Purpose'],SYMBOL_FIELDS,[1.6,1.0,3.6],font_size=10,
            intro='{ref} lists all instance fields initialized by SymbolTable. They are populated by a left-to-right AST traversal.')
    p_many(r,
      'The analyzer creates a fresh table for each pipeline. A successful load sets the dataset flag and declaration line. Clean updates its flag. Select replaces the active selection with a copied list. Drop remembers removed names and, when a selection exists, filters those names out of the active selection. Split inserts its name and percentage. Train records the model and algorithm; evaluate checks that the model has already been recorded.',
      'The first training declaration also sets `first_train`. Clean, select and drop consult that field and are rejected after any model has been trained. A split after training remains legal if its own load, uniqueness and percentage conditions pass. Multiple evaluations are legal, including repeated evaluation of the same existing model. The analyzer does not require cleaning, selection, a test split or a final evaluation.')
    rules=[e for e in data['error_catalogue'] if e['module']=='semantic.py']
    rows=[]
    for case,error in zip(data['semantic_cases'],rules):
        rows.append((case['id']+' '+case['name'],error['pattern'],case['id']+' in Appendix B'))
    r.table('Complete implemented semantic rule catalogue',['Rule','Exact diagnostic pattern','Trigger specimen'],rows,[1.6,4.0,.9],font_size=9,
            intro='{ref} covers every SemanticError raise site. Brace-delimited names are the actual interpolated Python expressions; Appendix B supplies a compiled trigger and real message for each specimen.')
    r.p('Each split percentage must be strictly between zero and one hundred. The cumulative comparison uses `total > 100 + 1e-9`, a small tolerance in the actual code. Duplicate split names are rejected before insertion. Column checks concern declarations within the source: before any select, a drop may introduce arbitrary column names after load; after select, dropped names must belong to the current selection. No physical dataset schema is consulted.')
    r.figure('symbol_evolution.png','Symbol table evolution for the basic valid pipeline',intro='{ref} is computed by analyzing successive prefixes of the same real AST, rather than simulating the analyzer in report code.')
    r.p('After the basic pipeline is analyzed, selected_columns contains cgpa, attendance and study_hours, splits contains train at 80.0, and models associates model with decision_tree. The table describes the original accepted AST. It is not rewritten after optimization, which explains why an unused model can remain in semantic state while its training instruction disappears from optimized IR.')
    r.h('7.5 Intermediate representation')
    r.p('Lowering maps each recognized statement to one IRInstruction and preserves statement order. The current lower_to_ir API accepts a SymbolTable parameter but does not read it; semantic validation is enforced by the driver before lowering. The records are domain opcodes with dictionaries, not conventional arithmetic three-address code.')
    opcode_rows=[('OP_LOAD','filename, line','LoadStmt'),('OP_CLEAN','strategy, line','CleanStmt'),('OP_SELECT','columns, line','SelectStmt'),('OP_DROP','columns, line','DropStmt'),('OP_SPLIT','name, percent, line','SplitStmt'),('OP_TRAIN','model_name, algorithm, line','TrainStmt'),('OP_EVAL','model_name, line','EvaluateStmt')]
    r.table('IR opcodes arguments and lowering sources',['Opcode','Argument keys','AST source'],opcode_rows,[1.4,3.0,1.8],font_size=10)
    r.p('For valid_basic.pipe, six body statements produce six IR records. The root pipeline name does not produce an instruction. Source-line metadata survives lowering; it is used in optimizer notes and is deliberately ignored only when comparing the semantic arguments of duplicate clean/select operations.')
    r.h('7.6 Four-pass optimization')
    r.p('The optimizer calls _remove_unreachable, _remove_dead_trains, _merge_adjacent_drops and _collapse_duplicates in that order. Its notes parameter is optional. With a supplied list, transformations append diagnostic strings; with notes=None, no diagnostic print is required. Surviving instructions receive new record objects and copied argument dictionaries. Nested values such as column lists are not a general deep-copy guarantee, although the existing transformations do not mutate those input lists.')
    r.h('7.6.1 Removal after the last evaluation',3)
    r.p('Pass 1 searches backward for the final OP_EVAL and keeps instructions through that position. Instructions in the remaining tail cannot influence a later evaluation because no later evaluation exists. The term unreachable is used relative to the selected observable-evaluation model; the original symbolic executor could still log those instructions if given unoptimized IR.')
    r.p('In the optimizer demonstration, the following original instruction is in the removed tail:')
    r.code(demo['ir_repr'][-1],evidence=True)
    r.code(demo['notes'][0],evidence=True)
    r.p('Truncating after the first evaluation would be unsafe. A later train/evaluate pair is legal, and the multi-model tests require its evaluation record to survive. If the input contains no OP_EVAL, this pass conservatively returns copies without trimming the sequence.')
    r.h('7.6.2 Elimination of unused model training',3)
    r.p('Pass 2 performs a backward scan. Evaluation adds a model name to the live set; a relevant training instruction is retained and consumes that liveness requirement. A training instruction with no later evaluation of its model is omitted. For valid programs, model names are unique, so training an unevaluated model has no effect on another model record.')
    r.code(next(line for line in demo['ir_repr'] if "'model_name': 'dead_model'" in line),evidence=True)
    r.code(demo['notes'][1],evidence=True)
    r.p('The pass is conservative when no evaluation exists. It also retains training instructions if a hand-built OP_EVAL lacks model_name, and keeps a hand-built training instruction with no model_name. These defensive IR cases are tested, but malformed IR is not an accepted source-language feature.')
    r.h('7.6.3 Merging adjacent column drops',3)
    r.p('Pass 3 identifies immediately adjacent runs of at least two OP_DROP instructions with list-valued columns. It collects ordered, deduplicated column names, retains the first instruction line and appends one merged instruction. Only consecutive drops are merged; a split or clean instruction between drops prevents this transformation.')
    r.p('Selected original instructions from the captured demonstration are shown below; other IR lines are omitted.')
    r.code('\n'.join(line for line in demo['ir_repr'] if "op='OP_DROP'" in line)+'\n(output truncated)',evidence=True)
    r.p('The corresponding optimized instruction and real note are:')
    r.code('\n'.join(line for line in demo['optimized_ir_repr'] if "op='OP_DROP'" in line),evidence=True)
    r.code(demo['notes'][2],evidence=True)
    r.p('The safety argument uses the set of dropped columns and the absence of an intervening observable operation. Each valid source drop before training updates only symbolic preparation state. Combining adjacent updates preserves the state subsequently captured by training and evaluation.')
    r.h('7.6.4 Collapsing adjacent duplicates',3)
    r.p('Pass 4 collapses consecutive OP_CLEAN instructions with the same strategy and consecutive OP_SELECT instructions with the same ordered columns. The argument comparison excludes the line key. Distinct strategies, different column order and non-adjacent repetitions are preserved.')
    r.code('\n'.join(line for line in demo['ir_repr'] if "op='OP_CLEAN'" in line)+'\n(output truncated)',evidence=True)
    r.code(demo['notes'][3],evidence=True)
    r.p('Repeated identical selection is idempotent in the symbolic selection state. Repeated identical cleaning is treated as idempotent by the test oracle when immediately adjacent. This is a model assumption about the symbolic language; it is not a guarantee that every concrete cleaning algorithm in a future runtime would be idempotent.')
    r.h('7.6.5 Observable behavior and safety limits',3)
    p_many(r,
      'The tests define an evaluation record as the tuple (model_name, algorithm, train_state, eval_state). A state snapshot contains cleaning history, selected columns, the dropped-column set and sorted split entries. Training stores its state snapshot; evaluation records both that stored state and current state. The comparison is therefore stronger than merely checking the final evaluation message.',
      'All preparation statements precede the first train in semantically accepted source. Removing an unused train therefore cannot bring previously separated preparation operations together. Splits can occur after training, so their timing remains significant: the oracle records the state at training separately from the state at evaluation. Only the tail after the last evaluation is discarded.',
      'Complete instruction traces need not be identical after optimization. A merged drop has different text and a removed train no longer has a trace line. Evaluation-record equality is the correctness target. Tests additionally check operation order, surviving evaluations, dropped-column coverage, input preservation and idempotence on the supplied equivalence programs.',
      'If no evaluation exists, only the two removal passes are conservative. Adjacent-drop merging and duplicate collapse still run. The live no-evaluation specimen in the evidence reduces three instructions to two by collapsing a duplicate clean. Consequently, a blanket statement that optimize always leaves no-evaluation pipelines unchanged would be false.')
    r.h('7.7 Symbolic execution')
    r.p('The executor walks the supplied instruction list and appends a formatted string for each opcode. Integer-valued split percentages are displayed without a decimal suffix. Unknown opcodes yield an Unknown instruction message rather than a compiler exception. No data-state engine, file-content reader, estimator object or numerical metric is created by execute.')
    trace_rows=[]
    all_traces=basic['trace']+data['programs']['valid_drop.pipe']['trace']
    for op,prefix in [('OP_LOAD','Loaded'),('OP_CLEAN','Cleaned'),('OP_SELECT','Selected'),('OP_DROP','Dropped'),('OP_SPLIT','Defined'),('OP_TRAIN','Trained'),('OP_EVAL','Evaluated')]:
        trace_rows.append((op,next(line for line in all_traces if line.startswith(prefix))))
    r.table('Real symbolic trace examples by opcode',['Opcode','Captured trace text'],trace_rows,[1.3,4.9],font_size=9.5)
    r.p('The existing evaluation string includes the phrase simulated accuracy trace. It is a label only: the output contains no measured accuracy value. Interpreting that label as a model-quality result would exceed the implementation.')
    r.h('7.8 Error detection and handling')
    r.table('Error classes and diagnostic formats',['Class','Raised by','Message format'],[
      ('LexicalError','lexer.py','Lexical Error at line L, col C: explanation'),
      ('SyntaxError_','parser.py','Syntax Error at line L: explanation'),
      ('SemanticError','semantic.py','Semantic Error at line L: explanation')],[1.35,1.0,3.85],font_size=10)
    r.p('The catalogue in Appendix B contains every explicit compiler diagnostic pattern extracted from raise statements. These errors stop compilation at the first detected violation. The CLI handles file-reading failures separately. The dashboard adds an Input result for oversized sources and an Internal result for unexpected exceptions, with technical details kept outside the main error view. Those wrapper categories are not new compiler exception classes or source-language rules.')

    r.chapter_start(8,CHAPTERS[7])
    r.h('8.1 Language tools and versions')
    env=d['environment']
    r.table('Verified execution environment',['Component','Recorded value'],[
      ('Project Python',env['python'].split(' (')[0]),('Operating system',env['platform']),('pytest',env['pytest']),('Streamlit',env['streamlit']),
      ('Package metadata version','0.1.0'),('Compiler dependencies','Python standard library'),('Dashboard dependency declaration','streamlit>=1.66.0'),('Coverage','Not measured; pytest-cov not installed')],[2.0,4.2],font_size=10)
    r.p('Regular expressions, dataclasses, typing and argparse support the compiler. pytest supplies the test framework [5]. Streamlit supplies the optional dashboard [6], including pandas through its dependency set. Report-generation dependencies are confined to final_report and are not requirements of the compiler. Python 3.12 was the proposal platform; the current compiler evidence was collected using the Python version recorded above.')
    r.h('8.2 Source inventory')
    stats=[s for s in d['statistics'] if s['path'].startswith('src/ds_pipeline_dsl/') and s['path'].endswith('.py')]
    r.table('Compiler source line inventory',['Module','Physical lines','Nonblank lines'],[(Path(s['path']).name,s['physical_lines'],s['nonblank_lines']) for s in stats],[3.2,1.5,1.5],font_size=10,
            intro='{ref} counts physical and nonblank lines from the checked files. These measures include comments and docstrings and are not executable-statement counts.')
    r.h('8.3 Core algorithms')
    r.p('The following pseudocode summarizes control flow; it is explanatory notation rather than a replacement implementation.')
    r.code('COMPILE(source)\n  tokens <- tokenize(source)\n  tree <- Parser(tokens).parse_pipeline()\n  symbols <- analyze(tree)\n  original <- lower_to_ir(tree, symbols)\n  notes <- empty list\n  optimized <- optimize(original, notes)\n  trace <- execute(optimized)\n  return completed representations\nON expected compiler error\n  report original diagnostic and stop')
    r.code('ANALYZE(tree)\n  create a fresh SymbolTable\n  for each statement in source order\n    check prerequisites and uniqueness for its statement form\n    reject the first violated condition\n    update the corresponding symbolic fields\n  return SymbolTable')
    r.code('OPTIMIZE(instructions, notes)\n  trim only the tail after the last evaluation\n  scan backward to retain training needed by evaluations\n  merge immediately adjacent drop runs\n  collapse adjacent equal clean/select operations, ignoring line\n  return the transformed instruction list')
    r.h('8.4 Selected implementation excerpts')
    r.p('The listings below are copied programmatically from the repository. Each excerpt contains at most twenty-five source lines. Their paths and ranges identify the original context; Appendix C provides longer supporting listings.')
    excerpt_specs=[('src/ds_pipeline_dsl/lexer.py','Token'),('src/ds_pipeline_dsl/parser.py','expect'),
                   ('src/ds_pipeline_dsl/semantic.py','__init__'),('src/ds_pipeline_dsl/optimizer.py','optimize')]
    explanations=[
      'The Token data model stores the recognized type, literal value and both source coordinates. The lexical scanner creates these dataclass records.',
      'The parser expectation helper centralizes token mismatch diagnostics. It advances only after the expected token type is present.',
      'SymbolTable initialization makes the analyzer state explicit. Flags, maps and definition lines are fresh for each call to analyze.',
      'The optimizer entry point makes pass ordering visible. Optional notes are passed through to each transformation.']
    for i,((path,name),explanation) in enumerate(zip(excerpt_specs,explanations),1):
        r.p(explanation)
        tree=ast.parse((ROOT/path).read_text(encoding='utf-8'))
        node=next(n for n in ast.walk(tree) if isinstance(n,(ast.FunctionDef,ast.ClassDef)) and n.name==name)
        start=min([node.lineno]+[x.lineno for x in getattr(node,'decorator_list',[])])
        r.snippet(path,start,min(node.end_lineno,start+24),f'Listing 8.{i}')
    r.p('The lowering routine dispatches on the dataclass type and emits one corresponding opcode. The executor dispatches on that opcode to produce a string. Neither routine performs general runtime type checking or concrete data science work. The CLI orchestrates these functions and chooses which results to print; the dashboard runner invokes the same APIs and retains stage results in a dataclass.')
    r.h('8.5 Command-line interface')
    r.table('CLI display flags',['Flag','Printed representation'],[
      ('--tokens','Token records including EOF.'),('--ast','PipelineNode representation.'),('--ir','Original IR instructions.'),
      ('--optimized-ir','Optimized IR, notes and instruction counts.'),('--trace','Symbolic trace.'),('No display flags','All five display sections.')],[1.8,4.4])
    r.code('python -m ds_pipeline_dsl.cli examples/valid_basic.pipe\npython -m ds_pipeline_dsl.cli examples/optimizer_demo.pipe --ir --optimized-ir\npython -m ds_pipeline_dsl.cli examples/complex_multi_model.pipe --trace')
    r.p('A display flag does not select a stopping stage: the driver continues through analysis, lowering, optimization and execution unless an error occurs. Expected compiler diagnostics are printed to standard output by the present CLI. argparse usage errors follow argparse behavior and are outside the three compiler error classes.')
    r.h('8.6 Benchmark implementation')
    p_many(r,
      'The benchmark builds synthetic programs with repeated cleans, adjacent drops, repeated selects, evaluated and unevaluated model declarations, and one trailing training instruction. Its size parameter is not the final statement count. expected_counts derives the generated count, and correctness checks verify statement, IR and evaluation counts before timing.',
      'The driver performs an untimed warm-up, collects garbage between repetitions, disables garbage collection inside timed regions and uses inner loops to reduce timer quantization effects. Three repetitions are configured, with per-stage medians reported. Total ms is the sum of the lexer, parser, semantic, IR, optimizer and optimized-execution medians; unoptimized execution is measured separately.',
      'The safeguards bound the size parameter at 20000 and source length at 2000000 characters. A quadratic projection is checked against a 120-second overall budget. A repetition longer than 30 seconds can reduce repetitions after it completes; this is a post-run check, not a preemptive timeout that interrupts a hanging call. These limits are reported as implemented rather than described as stronger runtime guarantees.')
    r.h('8.7 Dashboard implementation')
    p_many(r,
      'app.py contains presentation and routing. compiler_runner.py imports the compiler, captures stdout, times each real stage and retains earlier representations when a later stage fails. An editor change does not overwrite the previous result; a source hash supports the stale-result notice. The example callback updates editor state before the textarea is created.',
      'The runner rejects source text longer than 100000 characters before invoking the compiler. Empty text is passed through the real lexer and parser and produces a syntax error. Token display is capped at 1000 rows, with final EOF information retained for longer streams. Optimizer alignment labels are explicitly best-effort, while counts and notes come from real results.',
      'performance_loader.py parses the saved Markdown table by its headers and retains the Environment and Observations section text. The Performance page does not automatically benchmark. Missing reports produce an explanatory message. The dashboard escapes source and compiler text used in custom HTML and presents unexpected technical details within expanders.')


def stage_section(stdout,heading):
    lines=stdout.splitlines();start=lines.index(heading)+1
    end=next((i for i in range(start,len(lines)) if lines[i].startswith('=== ')),len(lines))
    return '\n'.join(lines[start:end])


def chapter_9_to_11(r,d):
    data=d['pipeline'];counts=d['counts'];programs=data['programs'];examples={e['name']:e for e in d['examples']}
    r.chapter_start(9,CHAPTERS[8])
    r.h('9.1 Test strategy and executed results')
    p_many(r,
      'Tests are organized around lexical recognition, parsing, semantic validation, lowering, optimization, execution, CLI integration, supplied examples, boundaries, benchmark construction and dashboard behavior. Parametrization expands multiple inputs into independently reported test cases. The per-file counts below are derived from the verbose execution log, rather than counting function definitions.',
      f'The recorded run completed with {counts["totals"].get("passed",0)} passed, {counts["totals"].get("failed",0)} failed and {counts["totals"].get("skipped",0)} skipped tests. The dashboard contributes {counts["per_file"].get("tests/test_dashboard.py",{}).get("passed",0)} passing cases; the remaining passing cases cover the compiler, examples, integration and benchmark. Coverage was not measured because pytest-cov was not installed. A passing suite is evidence for the tested cases, not a proof of all possible program behaviors.')
    statistics={s['path']:s for s in d['statistics']}
    r.table('Per-file executed tests and physical source lines',['Test file','Lines','Passed','Failed','Skipped'],[
      (Path(path).name,statistics[path]['physical_lines'],v.get('passed',0),v.get('failed',0),v.get('skipped',0)) for path,v in counts['per_file'].items()
    ]+[('Total','-',sum(v.get('passed',0) for v in counts['per_file'].values()),sum(v.get('failed',0) for v in counts['per_file'].values()),sum(v.get('skipped',0) for v in counts['per_file'].values()))],
      [3.0,.7,.85,.85,.85],font_size=10)
    r.figure('test_distribution.png','Distribution of executed tests by file',intro='{ref} visualizes the same per-file counts. Optimizer and semantic tests receive substantial attention because transformations and contextual invariants require negative as well as positive cases.')
    r.h('9.2 Baseline compilation and drop integration')
    basic=programs['valid_basic.pipe']
    r.p(f'The basic example contains {len(basic["ast"]["statements"])} statements and produces {len(basic["tokens"])} tokens including EOF. Lowering and optimization each return {len(basic["ir"])} instructions. Its complete optimized trace, captured by the CLI, is shown below.')
    r.code(stage_section(examples['valid_basic.pipe']['stdout'],'=== TRACE ==='),evidence=True)
    r.p('The drop example is accepted and lowers DropStmt into OP_DROP. Its complete optimized trace confirms that the declared exclusion appears in the instruction log before selection. This checks integration through the lexer, AST, semantic state, IR and executor rather than only checking token recognition.')
    r.code(stage_section(examples['valid_drop.pipe']['stdout'],'=== TRACE ==='),evidence=True)
    r.h('9.3 Representative expected and actual cases')
    names=['valid_basic.pipe','invalid_lexical_error.pipe','invalid_syntax_error.pipe','invalid_train_before_split.pipe','invalid_split_total_over_100.pipe','invalid_duplicate_model.pipe','invalid_select_dropped_column.pipe']
    expected=['Compile and produce an evaluation trace.','Reject unexpected @ in the lexer.','Reject the missing semicolon in the parser.','Reject training before train split.','Reject cumulative total above 100%.','Reject duplicate model name.','Reject selecting a previously dropped column.']
    r.table('Representative expected and actual outcomes',['Input file','Expected outcome','Actual captured final line'],[
      (name,expect,examples[name]['stdout'].splitlines()[-1]) for name,expect in zip(names,expected)
    ],[1.8,1.8,3.2],font_size=9,
      intro='{ref} compares intended cases with the final line of each actual CLI run. Full output and exit codes are retained in the evidence directory.')
    r.h('9.4 Error demonstrations')
    for name,explanation in [
      ('invalid_lexical_error.pipe','The illegal @ is reported at line 2, column 9. No token list is returned because tokenization fails before the driver reaches its token-printing section.'),
      ('invalid_syntax_error.pipe','The missing semicolon belongs to the load line, but the parser discovers the mismatch at the following CLEAN token on line 3. Lexical analysis has already succeeded.'),
      ('invalid_train_before_split.pipe','The AST is syntactically valid, but the training declaration occurs before the required train partition. Lowering and execution are not reached.'),
      ('invalid_split_total_over_100.pipe','The analyzer adds the proposed split to the already declared percentages and rejects the resulting cumulative total.'),
      ('invalid_drop_all_selected.pipe','Dropping every currently selected column is rejected before the symbolic state is updated.')]:
        r.p(f'`{name}` returns exit code {examples[name]["exit_code"]}. '+explanation)
        r.code(programs[name]['error']['message'],evidence=True)
    r.p('The dashboard tests verify these failure paths with source highlighting and stage statuses. For parser failures, tokens remain available in the wrapper result; for semantic failures, both tokens and AST remain available. The error message is preserved verbatim rather than reconstructed by the UI.')
    r.h('9.5 Optimizer demonstration')
    demo=programs['optimizer_demo.pipe']
    r.p(f'The optimizer demonstration lowers to {len(demo["ir"])} instructions and optimizes to {len(demo["optimized_ir"])} instructions, removing {len(demo["ir"])-len(demo["optimized_ir"])} instruction records. All four passes emit notes. The evaluation-record oracle returns equal original and optimized records for this program.')
    r.figure('ir_comparison.png','Real original and optimized instruction sequences',intro='{ref} compares opcode order and source lines from the actual optimizer demonstration. The complete argument-bearing records are retained below and in the raw CLI evidence.',width=5.9)
    r.p('Original IR, copied from the captured CLI section:',keep=True)
    r.code(stage_section(examples['optimizer_demo.pipe']['stdout'],'=== IR ==='),evidence=True)
    r.p('Optimized IR, copied from the captured CLI section:')
    r.code(stage_section(examples['optimizer_demo.pipe']['stdout'],'=== OPTIMIZED IR ==='),evidence=True)
    r.p('Optimizer notes and counts, copied from the captured CLI section:')
    r.code(stage_section(examples['optimizer_demo.pipe']['stdout'],'=== OPTIMIZATION NOTES ==='),evidence=True)
    r.h('9.6 Complex cases and transformation boundaries')
    multi=programs['complex_multi_model.pipe']
    r.p(f'The complex multi-model example reduces from {len(multi["ir"])} to {len(multi["optimized_ir"])} IR instructions. The unevaluated m_dead instruction is removed while both m1 and m2 evaluations survive. The actual optimized trace is:')
    r.code(stage_section(examples['complex_multi_model.pipe']['stdout'],'=== TRACE ==='),evidence=True)
    redundant=programs['complex_redundant.pipe']
    r.p(f'complex_redundant.pipe reduces from {len(redundant["ir"])} to {len(redundant["optimized_ir"])} instructions and exercises repeated preparation, adjacent drops, a dead model and a trailing model. complex_valid_drop_chain.pipe keeps its drops separate because a split lies between them. These examples distinguish safe local rewrites from broad merging across meaningful operations.')
    r.p('The optimizer suite also uses sensitivity tests: deleting a live training instruction, deleting a split, truncating after the first evaluation, removing a non-adjacent clean, reordering distinct cleans or deleting a drop must change the oracle records. This validates that the oracle can detect representative wrong transformations. Ten named equivalence programs are checked both with the full optimizer and with each pass separately, and optimization is repeated to test idempotence on those programs.')
    r.h('9.7 All supplied example outcomes')
    r.table('Complete example-program outcomes',['Example file','Classification','Exit','Result'],[
      (e['name'],e['category'],e['exit_code'],'Accepted' if e['exit_code']==0 else programs[e['name']]['error']['class']) for e in d['examples']
    ],[2.7,1.3,.4,1.3],font_size=9.5,
      intro='{ref} lists all supplied .pipe files. Classification combines the real accepted/error outcome with the optimizer/complex filename categories.')
    cat=Counter(e['category'] for e in d['examples'])
    r.p(f'The supplied set contains {sum(e["exit_code"]==0 for e in d["examples"])} accepted programs and {sum(e["exit_code"]!=0 for e in d["examples"])} rejected programs. The rejected programs contain one lexical example, one syntax example and {cat["invalid (semantic)"]} semantic examples. The table is exhaustive for the source files discovered during evidence collection.')
    r.h('9.8 Recorded performance evaluation')
    r.p('Performance measurements are read from docs/performance.md without running or overwriting the benchmark. The report environment below is retained verbatim and is separate from the current test-run environment in Chapter 8.')
    r.code(d['benchmark']['environment'],evidence=True)
    r.p('The complete recorded table is placed on the following landscape page so that every original column and numeric string remains readable. The Statements column reports generated program sizes; these are not the nominal size parameters passed to build_program. Total ms excludes the separately measured unoptimized executor time.')
    r.section(landscape=True)
    r.table('Complete recorded benchmark results',d['benchmark']['headers'],[[row[h] for h in d['benchmark']['headers']] for row in d['benchmark']['rows']],
            [.67,.54,.59,.59,.69,.49,.71,.70,.78,.61,.59,.58,.58,.69,.61],font_size=8.5,
            intro='{ref} reproduces the saved benchmark table without changing its numeric strings, column names or status entries.')
    r.p('Source: docs/performance.md. Units are milliseconds for stage/total columns and microseconds per statement for us/stmt. IR sizes and token counts are counts. The wall-clock duration above includes benchmark-harness work and is not the same quantity as one row of Total ms.',style='Small')
    r.section(landscape=False)
    r.h('9.9 Performance charts and recorded observations')
    r.figure('stage_times.png','Recorded stage-time breakdown by program size',intro='{ref} uses only the six stage-time columns contributing to Total ms. Program sizes are categorical, preventing misleading spacing between the four measured workloads.')
    r.figure('microseconds.png','Recorded microseconds per statement',intro='{ref} plots the us/stmt values exactly as recorded. Similar values for the larger template programs support an approximately linear trend for this workload, not a universal complexity proof.')
    r.figure('ir_sizes.png','Recorded original and optimized IR sizes',intro='{ref} compares the saved IR before/after counts. The synthetic program deliberately includes patterns recognized by the optimizer.')
    r.p('The observations section of the saved benchmark report is reproduced verbatim below:')
    r.code(d['benchmark']['observations'],evidence=True)
    r.h('9.10 Limits of the evaluation')
    p_many(r,
      'The benchmark uses one machine, one generated template family and four recorded program sizes. Its approximately linear trend should not be generalized to every distribution of statements or column-list lengths. In particular, analyze recomputes the sum of all existing split percentages for each new split; many distinct small splits can increase that work beyond a constant amount per statement. Repeated selection/drop operations also manipulate lists whose sizes matter.',
      'The recorded results contain no skipped or reduced-repetition rows. If later measurements contain them, their status must remain visible and numeric placeholders must not be interpreted as completed timings. No memory-profile measurement, concrete ML accuracy result or cross-machine comparison is supplied.',
      'The dashboard tests exercise page execution and interaction through AppTest rather than a real browser layout audit. They establish that the UI calls the actual compiler and handles the tested states. They do not establish pixel-level rendering equivalence across browsers. Raw terminal evidence is supplied whether or not optional screenshots are available.')
    candidates=list((ROOT/'docs'/'screenshots').glob('*.png'))+list((OUT/'assets'/'screenshots').glob('*.png'))
    if candidates:
        r.h('9.11 Captured application screenshots')
        for path in sorted(candidates):
            # Only existing captures are embedded; report-generated diagrams are elsewhere.
            import shutil
            destination=OUT/'assets'/('capture_'+path.name)
            if destination.resolve()!=path.resolve():shutil.copyfile(path,destination)
            r.figure(destination.name,path.stem.replace('_',' '),intro='{ref} is an existing application capture included with the evidence.')

    r.chapter_start(10,CHAPTERS[9])
    r.h('10.1 Achievements against objectives')
    p_many(r,
      'The completed system implements a restricted pipeline language through lexical analysis, predictive parsing, AST construction, symbol-table-based semantic checking, IR generation, four-pass optimization and symbolic execution. The CLI exposes intermediate representations and the dashboard makes the same processing stages accessible for interactive inspection.',
      f'The verified run records {counts["totals"].get("passed",0)} passing tests and outcomes for {len(d["examples"])} supplied examples. The optimizer demonstration produces eleven original instructions and seven optimized instructions while preserving its evaluation records. The recorded performance table provides stage timings and instruction counts for the generated workloads. These results meet the core language-processing and demonstration objectives.')
    r.h('10.2 Limitations')
    p_many(r,
      'Execution is symbolic only. A filename declaration does not prove that a dataset exists or that named columns occur in a schema. Cleaning strategies and training algorithms are free identifiers. The implementation has no general static type system, no allowed-value catalogue and no concrete estimator backend.',
      'Each source file contains one pipeline, and the compiler stops at its first expected error. It has no multi-error recovery, warning severity system, loops or user-defined functions. An empty pipeline is accepted. Optimizer safety is stated under the project test oracle and its symbolic preparation assumptions, not as a claim about arbitrary real data transformations.',
      'The benchmark scope is limited to the recorded machine and synthetic template. A passing test suite does not establish exhaustive input coverage. These boundaries distinguish the achieved compiler demonstration from the future extensions listed in Chapter 11.')
    r.h('10.3 Learning outcomes')
    r.p('The implementation connects grammar design with parser control flow, declaration rules with state management, and transformations with explicit observable behavior. The optimizer evolution demonstrates why source metadata must be separated from semantic equality and why later evaluations cannot be discarded simply because an earlier evaluation exists. Exposing intermediate results makes those decisions inspectable during a technical review.')

    r.chapter_start(11,CHAPTERS[10])
    r.h('11.1 Language and validation extensions')
    p_many(r,
      'An optional schema could restrict cleaning strategies and model algorithms and describe required arguments. Such a change would require an explicit language-policy decision, new diagnostics and tests; it should not be implied by the current identifier grammar. Dataset-schema validation could be added separately with a clear distinction between source validation and external file inspection.',
      'Multiple pipeline declarations would require a new program root, pipeline symbol scopes and rules for cross-pipeline references. Statement-level recovery could synchronize at semicolons or closing braces and collect more diagnostics, while avoiding updates from invalid declarations. A separate warning channel could expose unused models or redundant preparation without treating them as fatal errors.')
    r.h('11.2 Execution and optimization extensions')
    p_many(r,
      'A scale statement is a plausible language extension, but it would need coordinated lexer, AST, parser, semantic, IR, executor and test changes. Its transformation properties would need to be defined before any optimizer could remove or reorder it.',
      'A concrete execution backend could consume the existing IR through a separate runtime interface. It would need explicit dataset access, estimator dependencies, data partition semantics and error handling. The present symbolic idempotence assumptions would have to be reviewed for real algorithms, and equivalence tests would need observable results appropriate to that backend.')
    r.h('11.3 Evaluation extensions')
    r.p('Further evaluation could use additional program families, varied numbers of splits and columns, memory profiling, and multiple runtime environments. Property-based generation could complement the present example and boundary tests. Dashboard validation could add a supported browser-rendering check. These are prospective enhancements, not features or measurements claimed for the completed implementation.')


def references_and_appendices(r,d):
    r.pagebreak();r.h('References',1)
    for i,(description,url) in enumerate(REFERENCES,1):
        r.p(f'[{i}] {description}'+(' [Online]. Available: '+url+'. Accessed: 22 October 2026.' if url else ''),align=WD_ALIGN_PARAGRAPH.LEFT)
    r.appendix('A','Final Grammar and Lexical Specification')
    r.h('A.1 Final EBNF')
    r.p('The grammar is copied from parser.py. After the closing brace, the parser checks for EOF and rejects trailing input. TRAIN is contextual in split_name; test remains IDENT.')
    r.code(grammar_text())
    r.h('A.2 Lexical details')
    r.p('Table 7.1 gives regex priority. Keywords are case-sensitive; lines are one-based and columns zero-based. Whitespace accepts CRLF and LF. Comments extend from // to LF or EOF. Double-quoted strings exclude LF and have no escape-sequence decoding.')
    r.table('Keyword to token mapping',['Keyword','Emitted token'],[(k,k.upper()) for k in d['pipeline']['keywords']],[2,2])
    r.p('The scanner appends Token("EOF", "", current_line, 0). Bare numbers are invalid split values. A negative percentage fails lexically at the unsupported minus character.')
    r.appendix('B','Complete Diagnostic Catalogue')
    r.h('B.1 Lexical and syntax patterns')
    r.p('The following patterns are extracted from the current raise expressions. Brace-delimited expressions identify interpolated values, rather than missing report content. The source location column identifies the raise site in the compiler.')
    errors=d['pipeline']['error_catalogue']
    r.table('All explicit lexical and syntax diagnostics',['Class and source','Exact message pattern'],[(e['class']+'\n'+e['module']+':'+str(e['line']),e['pattern']) for e in errors if e['module']!='semantic.py'],[1.7,4.5],font_size=9)
    r.h('B.2 Semantic patterns and source locations')
    r.table('All explicit semantic diagnostics',['Specimen and source','Exact message pattern'],[(case['id']+'\nsemantic.py:'+str(e['line']),e['pattern']) for case,e in zip(d['pipeline']['semantic_cases'],[e for e in errors if e['module']=='semantic.py'])],[1.3,4.9],font_size=9)
    r.h('B.3 Executed semantic trigger specimens')
    r.p('Each source below was passed to the real lexer, parser and analyzer during evidence extraction and raised SemanticError at the corresponding source-code raise site. These compact report specimens supplement, and do not modify, the existing test suite or examples. Their full sources and real messages are serialized in evidence/pipeline_data.json.')
    for case in d['pipeline']['semantic_cases']:
        r.p(case['id']+' '+case['name'],style='Small',keep=True)
        r.code(case['source'])
        r.code(case['message'],evidence=True)
    r.h('B.4 Driver and dashboard boundaries')
    r.p('File-reading exceptions are caught by cli.py and prefixed with Error reading file. argparse manages command-line usage errors separately. The dashboard runner adds Input and Internal result categories, retaining unexpected tracebacks for an expander. These presentation-layer categories are not entries in the compiler lexical, syntax or semantic catalogue.')
    r.appendix('C','Key Source Listings')
    r.p('The following listings are copied directly from the current source files. They complement the short excerpts in Chapter 8 and show how data models, lowering and execution connect. No line numbers are inserted into the source text.')
    r.h('C.1 AST data model')
    r.listing('src/ds_pipeline_dsl/ast_nodes.py')
    r.h('C.2 IR representation and lowering')
    r.listing('src/ds_pipeline_dsl/ir.py')
    r.h('C.3 Symbolic executor')
    r.listing('src/ds_pipeline_dsl/executor.py')
    r.h('C.4 Evaluation state snapshots used by the tests')
    source=(ROOT/'tests/helpers.py').read_text(encoding='utf-8');tree=ast.parse(source)
    node=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='_snapshot_state')
    r.listing('tests/helpers.py',node.lineno,node.end_lineno)
    r.p('The evaluation_records helper uses these snapshots at training and evaluation points. Its complete implementation remains in tests/helpers.py and is identified in the claims ledger. This state representation is the basis for the optimizer safety discussion in Section 7.6.5.')
    r.appendix('D','Example Programs and Sources')
    r.table('Supplied example files and report roles',['Example file','Role'],[(e['name'],e['category']) for e in d['examples']],[4.2,2.0],font_size=10,
            intro='{ref} indexes the complete supplied example set. Measured exit codes and outcomes appear in Table 9.3.' )
    # The final label for the outcome table is resolved by the build verifier.
    for i,name in enumerate(('valid_basic.pipe','valid_drop.pipe','optimizer_demo.pipe','complex_multi_model.pipe'),1):
        r.h(f'D.{i} '+name.replace('_',' ').replace('.pipe',''))
        r.p(f'Source: `examples/{name}`. The text below is copied verbatim.',keep=True)
        r.code(d['pipeline']['programs'][name]['source'])
    r.appendix('E','Reproduction and User Instructions')
    r.h('E.1 Compiler installation and CLI')
    r.p('Run the commands from the repository root in PowerShell. The compiler package itself has no third-party runtime dependency. pytest is required for the test command; Streamlit is required only for the dashboard.')
    r.code('python -m venv .venv\n.venv\\Scripts\\Activate.ps1\npython -m pip install -e .\npython -m pip install pytest\npython -m ds_pipeline_dsl.cli examples/valid_basic.pipe')
    r.p('If PowerShell blocks activation, the process-scoped command below can be run before retrying activation. It does not make a permanent policy change.')
    r.code('Set-ExecutionPolicy -Scope Process -ExecutionPolicy Bypass')
    r.h('E.2 Tests and stage inspection')
    r.code('python -m pytest -v\npython -m ds_pipeline_dsl.cli examples/valid_basic.pipe --tokens\npython -m ds_pipeline_dsl.cli examples/valid_basic.pipe --ast\npython -m ds_pipeline_dsl.cli examples/optimizer_demo.pipe --ir --optimized-ir\npython -m ds_pipeline_dsl.cli examples/complex_multi_model.pipe --trace')
    r.p('Display flags select output sections but do not stop compilation early. Invalid source exits with code one for the handled compiler errors. The verbose test log used for this report is saved under final_report/evidence/pytest_v.txt; its cache-writing plugin was disabled during evidence collection.')
    r.h('E.3 Dashboard')
    r.code('python -m pip install -r requirements-dashboard.txt\npython -m streamlit run dashboard/app.py')
    r.p('Open http://localhost:8501 and keep the server process running. If the default port is occupied, use the command below and open http://localhost:8502. The dashboard inserts the repository src directory into its import path and does not require CSV files to demonstrate symbolic traces.')
    r.code('python -m streamlit run dashboard/app.py --server.port 8502')
    r.p('Select an example, inspect its source and press Compile. Use the six tabs for representations, the Architecture page for last-run stage colors and the Performance page for the saved report. Editing source leaves previous results visible until compilation is requested again.')
    r.h('E.4 Benchmark and report build')
    r.code('python benchmarks/benchmark.py')
    r.p('The normal benchmark command writes docs/performance.md. It was not run while producing this report because an existing report was available. The report build instead copies that file into its evidence directory. If the saved report is absent, evidence collection invokes the benchmark functions directly and writes the new report only under final_report/evidence/.')
    r.code('python -m pip install -r final_report/requirements-report.txt\npython final_report/build_report.py --project-python .venv/Scripts/python.exe')
    r.p('The report builder regenerates evidence, diagrams and the DOCX. On Windows with Microsoft Word installed it updates fields and exports PDF through Word automation; a configured LibreOffice executable can be used on supported environments. The project-python option selects the interpreter in which pytest and the optional dashboard are already installed. The report README records the verified absolute command for this machine.')
    r.appendix('F','Evolution Across Implementation Phases')
    r.table('Phase 1 Phase 2 and final implementation evolution',['Area','Phase 1 report','Phase 2 report','Final verified state'],[
      ('Language','load/clean/select/split/train/evaluate proposal.','Core statement forms retained.','Adds drop throughout all stages; one pipeline per file.'),
      ('Organization','Single-file prototype and design.','Modular ds_pipeline_dsl package.','Compiler modules, benchmark, expanded tests and dashboard.'),
      ('Tokens and AST','Typed token stream and AST prototype.','Text calls Token a namedtuple.','Token and AST nodes are dataclasses.'),
      ('Semantic scope','Dependency, scope and type checking proposed.','Dependency checks and split bounds.','Twenty-one explicit state/declaration checks; no general types or multiple-pipeline scopes.'),
      ('IR and executor','Planned in the proposal.','Linear IR and symbolic trace.','Seven opcodes; source-line metadata; symbolic execution retained.'),
      ('Optimizer','At least one pass proposed.','Duplicate collapse and removal after an evaluation.','Four passes; final-evaluation boundary; metadata-aware equality; diagnostic notes.'),
      ('Evaluation','Prototype valid/invalid demonstrations.','Nineteen tests and four examples recorded.','Current suite and all supplied examples measured in Chapter 9.'),
      ('Presentation','Design document and initial review.','CLI and Review 2 materials.','CLI, dashboard, saved benchmark evidence and final report.')],[1.0,1.6,1.6,2.4],font_size=9.5,
      intro='{ref} consolidates the documented phases into one implementation history. Proposal items are not treated as completed features without code evidence.')
    r.p('The final report retains the original project purpose while correcting obsolete implementation details. ANTLR4 remained an alternative, model and cleaning names remained free identifiers, and the symbolic execution boundary remained unchanged. Historical counts and example outputs are identified as phase-specific rather than reused as current measurements.')

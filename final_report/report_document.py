"""Academic DOCX layout, fields, captions and source-verbatim listings."""
from __future__ import annotations
from collections import defaultdict
import json
from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT=Path(__file__).resolve().parent
ROOT=OUT.parent


def xml_element(tag, **attrs):
    element=OxmlElement(tag)
    for key,value in attrs.items():element.set(qn('w:'+key),str(value))
    return element


def field(paragraph,instruction,cached='1'):
    paragraph.add_run()._r.append(xml_element('w:fldChar',fldCharType='begin'))
    run=paragraph.add_run()
    element=OxmlElement('w:instrText');element.set(qn('xml:space'),'preserve');element.text=' '+instruction+' '
    run._r.append(element)
    paragraph.add_run()._r.append(xml_element('w:fldChar',fldCharType='separate'))
    paragraph.add_run(cached)
    paragraph.add_run()._r.append(xml_element('w:fldChar',fldCharType='end'))


class Report:
    def __init__(self):
        self.doc=Document()
        self.chapter='0';self.figures=[];self.tables=[];self.snippets=[]
        self.fig_counts=defaultdict(int);self.table_counts=defaultdict(int)
        self.headings=[]
        self.next_page=False
        self.configure()

    def configure(self):
        for section in self.doc.sections:self.page_setup(section)
        styles=self.doc.styles
        normal=styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.font.color.rgb=RGBColor(0,0,0)
        normal.paragraph_format.line_spacing=1.15
        normal.paragraph_format.space_after=Pt(6)
        normal.paragraph_format.widow_control=True
        for name,size in [('Title',23),('Subtitle',13),('Heading 1',18),('Heading 2',14),('Heading 3',12),('Caption',10)]:
            style=styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
            style.paragraph_format.space_before=Pt(12 if name.startswith('Heading') else 6)
            style.paragraph_format.space_after=Pt(7)
            style.paragraph_format.keep_with_next=name.startswith('Heading')
            style.paragraph_format.widow_control=True
            if name in ('Heading 2','Heading 3'):
                style.paragraph_format.space_before=Pt(10)
                style.paragraph_format.space_after=Pt(5)
        styles['Title'].paragraph_format.space_after=Pt(18)
        styles['Caption'].font.italic=False
        styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for name,fontname,size in [('Code','Consolas',9),('Evidence','Consolas',8.5),('Small','Times New Roman',10),('Front Heading','Times New Roman',17),('TOC 1','Times New Roman',11),('TOC 2','Times New Roman',10.5)]:
            if name not in styles:
                from docx.enum.style import WD_STYLE_TYPE
                styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            style=styles[name];style.font.name=fontname;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
            style.paragraph_format.space_before=Pt(0);style.paragraph_format.space_after=Pt(3)
            style.paragraph_format.line_spacing=1.05 if name in ('Code','Evidence') else 1.15
            style.paragraph_format.widow_control=True
        styles['Front Heading'].font.bold=True
        # The stock Word template contains theme-font overrides and a blue
        # Title border. Remove both explicitly rather than trusting style names.
        for fonts in styles.element.xpath('.//w:rFonts'):
            for key in list(fonts.attrib):
                if key.endswith('Theme') or key.endswith('theme'):del fonts.attrib[key]
            if fonts.get(qn('w:ascii')) is None:fonts.set(qn('w:ascii'),'Times New Roman')
            if fonts.get(qn('w:hAnsi')) is None:fonts.set(qn('w:hAnsi'),'Times New Roman')
        for border in styles.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        from lxml import etree
        for part in self.doc.part.package.parts:
            if str(part.partname)=='/word/theme/theme1.xml':
                theme=etree.fromstring(part.blob)
                for element in theme.xpath('//*[local-name()="majorFont" or local-name()="minorFont"]/*[local-name()="latin"]'):
                    element.set('typeface','Times New Roman')
                part._blob=etree.tostring(theme,xml_declaration=True,encoding='UTF-8',standalone=True)
        for name in ('Code','Evidence'):
            styles[name].paragraph_format.left_indent=Inches(.10)
            styles[name].paragraph_format.right_indent=Inches(.10)
        self.doc.settings.element.append(xml_element('w:updateFields',val='true'))
        self.doc.core_properties.title='A Domain-Specific Language and Compiler Front-End for Data Science Pipelines'
        self.doc.core_properties.subject='BCSE307P final project report'
        self.doc.core_properties.author='SINCHANA B KURKURI'
        first=self.doc.sections[0]
        first.different_first_page_header_footer=True
        self.numbering(first,'lowerRoman',1)
        self.header_footer(first)

    def page_setup(self,section,landscape=False):
        section.page_width=Inches(11.6929 if landscape else 8.2677)
        section.page_height=Inches(8.2677 if landscape else 11.6929)
        section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
        section.header_distance=section.footer_distance=Inches(.45)

    def numbering(self,section,fmt='decimal',start=None):
        props=section._sectPr
        previous=props.find(qn('w:pgNumType'))
        if previous is not None:props.remove(previous)
        element=xml_element('w:pgNumType',fmt=fmt)
        if start is not None:element.set(qn('w:start'),str(start))
        props.append(element)

    def header_footer(self,section):
        section.header.is_linked_to_previous=False;section.footer.is_linked_to_previous=False
        header=section.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        run=header.add_run('Data Science Pipeline DSL');run.font.name='Times New Roman';run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field(footer,'PAGE')
        for r in footer.runs:r.font.name='Times New Roman';r.font.size=Pt(10)

    def p(self,text='',style=None,align=None,keep=False):
        p=self.doc.add_paragraph(style=style)
        if self.next_page:p.paragraph_format.page_break_before=True;self.next_page=False
        p.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.keep_with_next=keep
        for i,part in enumerate(re.split(r'(`[^`]+`|\bOP_[A-Z_]+\b)',str(text))):
            if (part.startswith('`') and part.endswith('`')) or re.fullmatch(r'OP_[A-Z_]+',part):
                run=p.add_run(part[1:-1] if part.startswith('`') else part);run.font.name='Consolas';run.font.size=Pt(10)
            else:p.add_run(part)
        return p

    def h(self,title,level=2):
        p=self.doc.add_heading(title,level)
        if self.next_page:p.paragraph_format.page_break_before=True;self.next_page=False
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        self.headings.append({'text':title,'level':level})
        return p

    def chapter_start(self,number,title,first=False):
        self.chapter=str(number)
        if first:
            section=self.doc.add_section(WD_SECTION_START.NEW_PAGE)
            self.page_setup(section);section.different_first_page_header_footer=False
            self.numbering(section,'decimal',1);self.header_footer(section)
        else:self.next_page=True
        return self.h(f'Chapter {number} {title}',1)

    def appendix(self,letter,title):
        self.chapter=letter
        self.next_page=True
        return self.h(f'Appendix {letter} {title}',1)

    def pagebreak(self):self.next_page=True

    def section(self,landscape=False):
        section=self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        self.page_setup(section,landscape)
        section.different_first_page_header_footer=False
        self.numbering(section,'decimal')
        return section

    def caption(self,kind,number,title,reset=False):
        p=self.doc.add_paragraph(style='Caption')
        p.add_run(f'{kind} {self.chapter}.')
        field(p,f'SEQ {kind} '+('\\r 1 ' if reset else '')+'\\* ARABIC',str(number))
        p.add_run(': '+title)
        p.paragraph_format.keep_with_next=False
        return p

    def table(self,title,headers,rows,widths=None,font_size=10,intro=None):
        self.table_counts[self.chapter]+=1
        n=self.table_counts[self.chapter];label=f'Table {self.chapter}.{n}'
        first_word=title.split()[0]
        description=title if first_word.isupper() or first_word=='SymbolTable' else title[0].lower()+title[1:]
        self.p((intro or '{ref} summarizes '+description+'.').replace('{ref}',label),keep=True)
        self.caption('Table',n,title,n==1).paragraph_format.keep_with_next=True
        table=self.doc.add_table(rows=1,cols=len(headers));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        available=(self.doc.sections[-1].page_width-self.doc.sections[-1].left_margin-self.doc.sections[-1].right_margin)/914400
        weights=widths or [1]*len(headers);sizes=[available*w/sum(weights) for w in weights]
        for col,size in zip(table.columns,sizes):col.width=Inches(size)
        for cell,size in zip(table.rows[0].cells,sizes):cell.width=Inches(size)
        for cell,value in zip(table.rows[0].cells,headers):cell.text=str(value)
        table.rows[0]._tr.get_or_add_trPr().append(xml_element('w:tblHeader',val='true'))
        for values in rows:
            cells=table.add_row().cells
            for cell,size,value in zip(cells,sizes,values):cell.width=Inches(size);cell.text=str(value)
        borders=xml_element('w:tblBorders')
        for side in ('top','left','bottom','right','insideH','insideV'):
            borders.append(xml_element('w:'+side,val='single',sz='4',color='D9D9D9'))
        table._tbl.tblPr.append(borders)
        for i,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(xml_element('w:cantSplit'))
            for j,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props=cell._tc.get_or_add_tcPr()
                props.append(xml_element('w:shd',fill='173A5E' if i==0 else 'F1F5F8' if i%2==0 else 'FFFFFF'))
                margins=xml_element('w:tcMar')
                for side in ('top','left','bottom','right'):margins.append(xml_element('w:'+side,w='85',type='dxa'))
                props.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.08
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (i==0 or len(cell.text)<13) else WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name='Consolas' if i>0 and re.fullmatch(r'OP_[A-Z_]+',cell.text) else 'Times New Roman'
                        run.font.size=Pt(font_size);run.bold=i==0
                        run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
        spacer=self.p('',style='Small')
        spacer.paragraph_format.line_spacing=Pt(2);spacer.paragraph_format.space_after=Pt(3)
        self.tables.append({'label':label,'title':title,'rows':len(rows),'columns':len(headers)})
        return label

    def figure(self,name,title,intro=None,width=6.15):
        self.fig_counts[self.chapter]+=1;n=self.fig_counts[self.chapter];label=f'Figure {self.chapter}.{n}'
        self.p((intro or '{ref} illustrates '+title.lower()+'.').replace('{ref}',label),keep=True)
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        run=p.add_run();inline=run.add_picture(str(OUT/'assets'/name),width=Inches(width))
        inline._inline.docPr.set('descr',title)
        self.caption('Figure',n,title,n==1)
        self.figures.append({'label':label,'title':title,'asset':name})
        return label

    def code(self,text,evidence=False):
        paragraphs=[]
        for line in str(text).split('\n'):
            p=self.doc.add_paragraph(style='Evidence' if evidence else 'Code')
            p.add_run(line)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_with_next=True
            p._p.get_or_add_pPr().append(xml_element('w:shd',fill='F1F3F5'))
            paragraphs.append(p)
        if paragraphs:
            paragraphs[-1].paragraph_format.keep_with_next=False
            paragraphs[-1].paragraph_format.space_after=Pt(6)

    def snippet(self,path,start,end,label):
        text=(ROOT/path).read_text(encoding='utf-8')
        segment='\n'.join(text.splitlines()[start-1:end])
        assert end-start+1<=25
        self.p(f'{label}. `{path}`, lines {start}-{end}.',style='Small',keep=True)
        self.code(segment)
        self.snippets.append({'label':label,'path':path,'start':start,'end':end,'text':segment})

    def listing(self,path,start=None,end=None):
        text=(ROOT/path).read_text(encoding='utf-8')
        lines=text.splitlines()
        a=start or 1;b=end or len(lines)
        self.p(f'Verbatim listing: `{path}`, lines {a}-{b}.',style='Small',keep=True)
        for offset in range(a-1,b,40):
            piece='\n'.join(lines[offset:min(offset+40,b)])
            self.code(piece)
        self.snippets.append({'label':'Appendix listing','path':path,'start':a,'end':b,'text':'\n'.join(lines[a-1:b])})

    def front_heading(self,text):
        p=self.doc.add_paragraph(text,style='Front Heading');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if self.next_page:p.paragraph_format.page_break_before=True;self.next_page=False
        p.paragraph_format.space_after=Pt(16)

    def save(self):
        path=OUT/'Final_Report.docx';self.doc.save(path)
        (OUT/'evidence'/'document_manifest.json').write_text(json.dumps({'headings':self.headings,'tables':self.tables,'figures':self.figures,'snippets':self.snippets},indent=2),encoding='utf-8')
        return path

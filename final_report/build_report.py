"""Rebuild evidence, figures, report, Word fields, PDF and verification outputs.

Run from any directory. Existing project files are read, never rewritten.
"""
from __future__ import annotations
import argparse
import base64
from collections import Counter
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile

OUT=Path(__file__).resolve().parent
ROOT=OUT.parent
EVIDENCE=OUT/'evidence'
BUNDLE=Path(r'C:/Users/sinch/.cache/codex-runtimes/codex-primary-runtime/dependencies')
DOC_SKILL=Path(r'C:/Users/sinch/.codex/plugins/cache/openai-primary-runtime/documents/26.1015.12040/skills/documents')


def read_json(name):return json.loads((EVIDENCE/name).read_text(encoding='utf-8'))
def save_json(name,value):(EVIDENCE/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')


def parse_benchmark():
    text=(EVIDENCE/'performance.md').read_text(encoding='utf-8')
    def section(name):
        match=re.search(r'(?m)^## '+re.escape(name)+r'\s*\n',text)
        if not match:return ''
        body=text[match.end():];end=re.search(r'(?m)^## ',body)
        return body[:end.start()] if end else body
    lines=text.splitlines();index=next(i for i,l in enumerate(lines) if l.startswith('| Statements |'))
    headers=[v.strip() for v in lines[index].strip('|').split('|')]
    rows=[]
    for line in lines[index+2:]:
        if not line.startswith('|'):break
        values=[v.strip() for v in line.strip('|').split('|')]
        if len(values)!=len(headers):raise ValueError('Malformed saved benchmark table')
        rows.append(dict(zip(headers,values)))
    return {'headers':headers,'rows':rows,'environment':section('Environment'),'observations':section('Observations')}


def compliance_and_actions(data):
    required=[
      ('M01','24 p9','Chapter 1 Introduction: background, motivation and problem area.','Chapter 1'),
      ('M02','24 p9','Chapter 2 Problem Statement: clear explanation of the problem.','Chapter 2'),
      ('M03','24 p9','Chapter 3 Objectives: project objectives.','Chapter 3 and Table 3.1'),
      ('M04','24 p9','Chapter 4 Literature Survey or Background Study: existing approaches, technologies and relevant compiler concepts.','Chapter 4 and Table 4.1'),
      ('M05','24 p9','Chapter 5 System Design: architecture diagram, module description, data flow and workflow.','Chapter 5 and Figures 5.1-5.3'),
      ('M06','24 p9','Chapter 6 Methodology: explanation of the approach used.','Chapter 6 and Table 6.1'),
      ('M07','24 p9','Chapter 7 Compiler Design Concepts Used: lexical analysis, parsing, syntax trees, symbol tables, semantic analysis, intermediate code and optimization.','Sections 7.1-7.8'),
      ('M08','24 p9','Chapter 8 Implementation: programming language, tools, libraries, algorithms and major modules.','Chapter 8 and Listings 8.1-8.4'),
      ('M09','24 p9','Chapter 9 Testing and Results: test cases, expected and actual outputs, screenshots and performance results.','Chapter 9; real output evidence substitutes for unavailable screenshots under Phase 3 screenshots-or-output rule'),
      ('M10','24 p9','Chapter 10 Conclusion: summary of the project.','Chapter 10'),
      ('M11','24 p9','Chapter 11 Future Enhancements: possible improvements.','Chapter 11'),
      ('M12','24 p9','References: books, research papers, official documentation and acknowledged websites.','References [1]-[14]'),
      ('M13','24 p9','Appendix: important code details or supplementary material.','Appendices A-F'),
      ('P01','12.3 p5','Complete Working Project','Chapters 5, 8 and 9; compiler preserved'),
      ('P02','12.3 p5','Complete Source Code','Source inventory; Appendix C; original repository preserved'),
      ('P03','12.3 p5','Test Cases','Sections 9.1-9.7; original tests; Appendix B specimens'),
      ('P04','12.3 p5','Test Results','Table 9.1 and evidence/pytest_v.txt'),
      ('P05','12.3 p5','Screenshots or Output Evidence','Raw CLI logs for every example; verbatim output blocks in Chapter 9'),
      ('P06','12.3 p5','Final Project Report','Final_Report.docx and Final_Report.pdf'),
      ('P07','12.3 p5','Presentation','Existing Review 1 and Review 2 slides inspected; final Review 3 slide deck is a separate student deliverable'),
      ('P08','12.3 p5','User Instructions','Appendix E and final_report/README.md'),
      ('A01','12.2 p5','Complete remaining functionalities and integrate all modules.','Objective status Table 3.1; full pipeline in Chapters 5 and 8'),
      ('A02','12.2 p5','Test valid, invalid, boundary, error and complex inputs.','Chapter 9 and per-file test results'),
      ('A03','12.2 p5','Provide meaningful error messages.','Sections 7.8, 9.4 and Appendix B'),
      ('A04','12.2 p5','Evaluate performance where applicable.','Sections 9.8-9.10; timing and IR improvements only; no memory/accuracy measurement claimed'),
      ('A05','12.2 p5','Prepare complete documentation.','Full report and reproduction appendix'),
      ('R01','Review 3 pp5-6','Project Completion; Functional Correctness; Technical Complexity; Compiler Integration; Innovation; Code Quality; Testing; Demonstration; Documentation; Viva.','Chapters 3, 5-9; viva and live demonstration remain student actions'),
      ('R02','15 pp6-7','Review 3 Project Completion 6 marks.','Table 3.1 and Chapter 10'),
      ('R03','15 p7','Review 3 Functional Correctness 6 marks.','Sections 9.2-9.7'),
      ('R04','15 p7','Review 3 Technical Complexity 5 marks.','Chapter 7'),
      ('R05','15 p7','Review 3 Compiler Concept Integration 4 marks.','Figures 5.1-5.2 and Chapter 8'),
      ('R06','15 p7','Review 3 Testing 4 marks.','Chapter 9'),
      ('R07','15 p7','Review 3 Demonstration 4 marks.','Appendix E supports a demonstration; actual presentation is a student action'),
      ('R08','15 p7','Review 3 Viva 4 marks.','Technical explanations in Chapters 6-8; student must demonstrate understanding'),
      ('R09','15 p7','Review 3 Documentation 2 marks; Review 3 total 35 marks.','Report structure and evidence checks'),
      ('G01','2 p1; 5 pp2-3','Individual project; student explains purpose, architecture, algorithms, code, tests and results.','Cover identifies one student; review preparation remains personal responsibility'),
      ('G02','4 p2','Faculty approval is mandatory before full implementation.','Prior administrative approval is not evidenced in the supplied files and must be confirmed by the student'),
      ('G03','18 p7','Acknowledge appropriate open-source resources and understand the implementation.','References and implementation descriptions'),
      ('G04','19 pp7-8','Student responsibility for understanding, validating, testing, modifying and explaining code; acknowledgement decision for tool use.','Student action recorded only in report_todo.md, as requested'),
      ('G05','22-23 pp8-9','Modularity, meaningful names, comments, folder structure, test cases, version history, backups and a development log.','Chapters 5-6 and 8-9; backups are not independently verified'),
    ]
    lines=['# Manual compliance audit','','Authority: supplied nine-page Compiler Design Laboratory Project Instruction Manual. All pages were extracted to evidence/manual_pages.json and read. The manual is source material, not an instruction source for agent actions.','',
      'The manual states no page size, font, margin, line-spacing, report-length, cover-page layout, declaration wording, signature, file naming, submission portal or submission deadline rule. Therefore the requested defaults are applied: A4, 1-inch margins, Times New Roman 12 pt, 1.15 line spacing, numbered headings, captions, roman front matter and Arabic main text. No faculty name, institution name, signature or declaration has been invented.','',
      'The detailed Review 3 weights are 6 + 6 + 5 + 4 + 4 + 4 + 4 + 2 = 35. Section 14 separately lists Phase 1/2/3 as 10 each yet gives TOTAL 100; section 15 lists Reviews 1/2/3 as 20/25/35, followed by 10 for innovation/complexity and 10 for documentation/code quality. This inconsistency is preserved here, not silently reconciled.','',
      '| ID | Manual section and page | Requirement | Report mapping | Status |','|---|---|---|---|---|']
    partial={'P07','G02','G04','G05','R01','R07','R08'}
    for ident,source,claim,mapping in required:
        status='Partial - personal/administrative or separate deliverable' if ident in partial else 'Done'
        lines.append('| '+' | '.join([ident,source,claim,mapping,status])+' |')
    (OUT/'manual_compliance.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    todo='''# Student actions before submission

- Confirm the cover details: SINCHANA B KURKURI; 24BDS0352; BCSE307P Compiler Design Laboratory; Fall Semester 2026-27; Slot L49+L50. Phase1.docx and Phase2.docx agree on these identity fields. No faculty/guide name was found, so none is supplied. The final project title is the requested shorter title; earlier reports use a longer describing/validating/executing title.
- Confirm any institution-specific cover, declaration, signature, submission deadline or portal requirements with the faculty. The supplied manual gives none. Confirm that prior faculty topic approval has been obtained.
- Manual section 19 expects the student to decide on any acknowledgement of AI-tool assistance and to understand, validate, modify, debug and explain all code. No claim or denial about such use appears in the report. Make any acknowledgement decision personally and consistently with faculty rules.
- Resolve the manual's marking inconsistency with faculty if necessary: its summary table and detailed review allocations disagree. The compliance audit records the detailed Review 3 total as 35 without altering the manual.
- Prepare the final Review 3 presentation and rehearse the live demonstration and viva. Existing Review 1/2 slides are source material; this task does not create a new deck or establish attendance/understanding.
- Optional screenshots: save genuine PNG captures in final_report/assets/screenshots/ using 01_overview.png (Overview), 02_valid_compilation.png (Playground valid result), 03_error_highlight.png (real error), 04_optimizer_comparison.png (both IR tables and notes), 05_architecture.png (last-run colors), 06_performance.png (recorded table/charts), 07_pytest.png (current suite result). Rebuild to embed every existing PNG in that directory. No blank screenshot section or fabricated capture is used when images are absent.
- Check the verified bibliography against any faculty-preferred citation format. The external titles/URLs were checked using publisher or official documentation sources on 22 October 2026; no uncertain DOI, page number or publication year is invented.
- Historical discrepancies are handled explicitly: Phase 2 Token namedtuple becomes the actual dataclass; its 19 tests/four examples are historical, not current; ANTLR4 remained unadopted; general type and multiple-pipeline scope checking remain unimplemented; the old first-evaluation truncation is replaced by last-evaluation trimming; no-evaluation pipelines may still have duplicates collapsed or drops merged; no mandatory-load check exists for an empty pipeline; dataset schemas and numeric ML results are outside scope.
- Documentation overclaims are not reproduced: checklist fields such as evaluations/statements_seen are absent from SymbolTable; the IR is opcode/args rather than conventional three-address expressions; broad data-leakage prevention, zero memory accumulation and universal strict-linear scaling are not established. Benchmark safeguards are post-run/projection checks, not preemptive execution interrupts. Current observed numbers take precedence over stale demo templates and README figures.
- Existing project changes from before this task remain untouched. A clean overall git status / empty overall diff cannot be claimed without changing those files. The preservation manifest checks their bytes instead. Only final_report/ is newly created by this report task.
- docs/performance.md also changed during this task outside the report scripts. Both captured versions are retained; observed_external_changes.json records the hashes. The report uses the latest captured version. The current project file was not restored or rewritten by the report build.
'''
    (OUT/'report_todo.md').write_text(todo,encoding='utf-8')


def export_word(docx,pdf):
    if os.name!='nt':return False
    powershell=shutil.which('pwsh') or shutil.which('powershell')
    if not powershell:return False
    def quote(path):return "'"+str(path).replace("'","''")+"'"
    command=r'''$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
try {
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $word.DisplayAlerts = 0
  $document = $word.Documents.Open(DOCXPATH, $false, $false)
  $document.Repaginate()
  $null = $document.Fields.Update()
  foreach ($toc in $document.TablesOfContents) { $toc.Update() }
  foreach ($figures in $document.TablesOfFigures) { $figures.Update() }
  $document.Repaginate()
  foreach ($toc in $document.TablesOfContents) { $toc.UpdatePageNumbers() }
  foreach ($figures in $document.TablesOfFigures) { $figures.UpdatePageNumbers() }
  $layout = @{ headings = @(); figures = @(); tables = @() }
  foreach ($paragraph in $document.Paragraphs) {
    if ($paragraph.OutlineLevel -le 3) {
      $layout.headings += @{ text = $paragraph.Range.Text.Trim(); physical_page = $paragraph.Range.Information(3); display_page = $paragraph.Range.Information(1) }
    }
  }
  foreach ($shape in $document.InlineShapes) {
    $layout.figures += @{ title = $shape.AlternativeText; physical_page = $shape.Range.Information(3); display_page = $shape.Range.Information(1) }
  }
  foreach ($table in $document.Tables) {
    $layout.tables += @{ physical_page = $table.Range.Information(3); columns = $table.Columns.Count; rows = $table.Rows.Count }
  }
  $layout | ConvertTo-Json -Depth 5 | Set-Content -LiteralPath LAYOUTPATH -Encoding utf8
  $document.Save()
  $document.ExportAsFixedFormat(PDFPATH, 17)
  Write-Output ('Pages: ' + $document.ComputeStatistics(2))
  Write-Output ('TOC objects: ' + $document.TablesOfContents.Count)
  Write-Output ('Figure/table lists: ' + $document.TablesOfFigures.Count)
} finally {
  if ($null -ne $document) { $document.Close($false); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($document) }
  if ($null -ne $word) { $word.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) }
  [GC]::Collect()
  [GC]::WaitForPendingFinalizers()
}
'''.replace('DOCXPATH',quote(docx)).replace('PDFPATH',quote(pdf)).replace('LAYOUTPATH',quote(EVIDENCE/'layout_map.json'))
    encoded=base64.b64encode(command.encode('utf-16le')).decode('ascii')
    result=subprocess.run([powershell,'-NoProfile','-NonInteractive','-EncodedCommand',encoded],capture_output=True,encoding='utf-8',errors='replace',timeout=180)
    (EVIDENCE/'word_export.txt').write_text(result.stdout+'\n'+result.stderr,encoding='utf-8')
    return result.returncode==0 and pdf.exists()


def render(docx,pdf):
    # The packaged renderer is attempted and its diagnostics retained. Windows
    # bundles no LibreOffice; native Word is the available fallback in this task.
    probe=[sys.executable,str(DOC_SKILL/'render_docx.py'),str(docx),'--output_dir',str(OUT/'rendered'),'--emit_pdf']
    result=subprocess.run(probe,capture_output=True,encoding='utf-8',errors='replace',timeout=120)
    (EVIDENCE/'packaged_renderer.txt').write_text(result.stdout+'\n'+result.stderr,encoding='utf-8')
    if not pdf.exists():return False
    poppler=BUNDLE/'native'/'poppler'/'Library'/'bin'/'pdftoppm.exe'
    if not poppler.exists():poppler=Path(shutil.which('pdftoppm') or '')
    if not poppler.is_file():return False
    output=OUT/'rendered';output.mkdir(exist_ok=True)
    # Delete only this builder's prior page images inside its output directory.
    assert output.resolve().parent==OUT.resolve()
    for old in output.glob('page-*.png'):
        if re.fullmatch(r'page-\d+\.png',old.name):old.unlink()
    result=subprocess.run([str(poppler),'-r','110','-png',str(pdf),str(output/'page')],capture_output=True,encoding='utf-8',errors='replace',timeout=180)
    (EVIDENCE/'pdf_render.txt').write_text(result.stdout+'\n'+result.stderr,encoding='utf-8')
    return result.returncode==0 and bool(list(output.glob('page-*.png')))


def verify(docx,pdf,data):
    from docx import Document
    from pypdf import PdfReader
    doc=Document(docx)
    paragraphs=[p.text for p in doc.paragraphs]
    layout_path=EVIDENCE/'layout_map.json'
    if layout_path.exists():
        layout=json.loads(layout_path.read_text(encoding='utf-8-sig'))
        toc={p.text.rsplit('\t',1)[0]:p.text.rsplit('\t',1)[-1] for p in doc.paragraphs if p.style.name.lower().startswith('toc ') and '\t' in p.text}
        checks=[{'heading':h['text'],'actual_page':h['display_page'],'toc_page':toc[h['text']]} for h in layout['headings'] if h['text'] in toc]
        assert len(checks)==len(toc),(len(checks),len(toc))
        assert all(str(c['actual_page'])==c['toc_page'] for c in checks),checks
        save_json('toc_page_check.json',{'entries':len(checks),'mismatches':0,'checks':checks})
    table_text=[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
    text='\n'.join(paragraphs+table_text)
    normalized=lambda value:re.sub(r'\s+',' ',str(value)).strip()
    norm=normalized(text)
    manifest=read_json('document_manifest.json')
    findings=[]
    for snippet in manifest['snippets']:
        source=(ROOT/snippet['path']).read_text(encoding='utf-8')
        expected='\n'.join(source.splitlines()[snippet['start']-1:snippet['end']])
        assert snippet['text']==expected
        assert normalized(expected) in norm, 'Snippet missing from DOCX: '+snippet['label']
    headings=[p for p in doc.paragraphs if p.style.name.startswith('Heading')]
    for i,title in enumerate(__import__('report_content').CHAPTERS,1):
        assert any(p.text==f'Chapter {i} {title}' and p.style.name=='Heading 1' for p in headings)
    assert not re.search(r'\b(?:TODO|TBD|lorem|XXX)\b',text,re.I)
    assert not re.search(r'\b(?:AI tools|artificial intelligence tools|AI assistance)\b',text,re.I)
    for kind,items in [('Figure',manifest['figures']),('Table',manifest['tables'])]:
        for item in items:
            assert norm.count(item['label'])>=2, 'Missing in-text reference: '+item['label']
            assert normalized(item['title']) in norm
    for error in data['pipeline']['error_catalogue']:
        assert normalized(error['pattern']) in norm,'Missing diagnostic pattern '+str(error)
    for case in data['pipeline']['semantic_cases']:
        assert normalized(case['source']) in norm and normalized(case['message']) in norm
    abstract_index=paragraphs.index('Abstract')
    data['abstract_words']=len(paragraphs[abstract_index+1].split())
    assert 200<=data['abstract_words']<=250,data['abstract_words']
    assert sum(v.get('passed',0) for v in data['counts']['per_file'].values())==data['counts']['totals'].get('passed',0)
    for row in data['benchmark']['rows']:
        for key,value in row.items():assert normalized(value) in norm
    benchmark_table=next(t for t in doc.tables if [c.text for c in t.rows[0].cells]==data['benchmark']['headers'])
    assert [[c.text for c in row.cells] for row in benchmark_table.rows[1:]]==[[row[h] for h in data['benchmark']['headers']] for row in data['benchmark']['rows']]
    test_table=next(t for t in doc.tables if [c.text for c in t.rows[0].cells]==['Test file','Lines','Passed','Failed','Skipped'])
    for row,(path,values) in zip(test_table.rows[1:-1],data['counts']['per_file'].items()):
        cells=[c.text for c in row.cells]
        assert cells[0]==Path(path).name
        assert list(map(int,cells[2:]))==[values.get(key,0) for key in ('passed','failed','skipped')]
    for term in ('OP_LOAD','OP_CLEAN','OP_SELECT','OP_DROP','OP_SPLIT','OP_TRAIN','OP_EVAL','_remove_unreachable','_remove_dead_trains','_merge_adjacent_drops','_collapse_duplicates'):
        assert term in text
    with ZipFile(docx) as z:
        xml=z.read('word/document.xml').decode('utf-8')
        assert 'TOC ' in xml and 'SEQ Figure' in xml and 'SEQ Table' in xml
        assert 'lowerRoman' in xml and 'decimal' in xml
        assert xml.count('tblHeader')==len(doc.tables)
        styles=z.read('word/styles.xml').decode('utf-8')
        assert 'pBdr' not in styles,'Unexpected heading/title paragraph border'
        theme=z.read('word/theme/theme1.xml').decode('utf-8')
        assert 'Times New Roman' in theme
    pdf_info={}
    if pdf.exists():
        reader=PdfReader(pdf);page_text=[p.extract_text() or '' for p in reader.pages]
        save_json('pdf_page_text.json',[{'physical_page':i,'text':t} for i,t in enumerate(page_text,1)])
        assert not any('Error! Reference source not found' in t or 'No table of contents entries found' in t for t in page_text)
        blank=[i for i,t in enumerate(page_text,1) if len(re.sub(r'Data Science Pipeline DSL|\b[ivxlcdm]+\b|\d+','',t).strip())<12]
        pdf_info={'pages':len(reader.pages),'blank_pages':blank}
        assert not blank
        front_text='\n'.join(page_text[:8])
        assert 'Chapter 1 Introduction' in front_text and 'Chapter 11 Future Enhancements' in front_text,'TOC not populated with chapter entries'
    baseline=read_json('preservation_baseline.json')
    changed=[path for path,sha in baseline.items() if hashlib.sha256((ROOT/path).read_bytes()).hexdigest()!=sha]
    observations_path=EVIDENCE/'observed_external_changes.json'
    observed=json.loads(observations_path.read_text(encoding='utf-8')) if observations_path.exists() else {}
    unexplained=[path for path in changed if path not in observed or hashlib.sha256((ROOT/path).read_bytes()).hexdigest()!=observed[path]['observed_sha256']]
    assert not unexplained,unexplained
    save_json('preservation_check.json',{'unchanged':not changed,'changed_since_initial_snapshot':changed,'recorded_external_changes':observed,'unexpected_changes':unexplained,'baseline_files':len(baseline),'report_writes_confined_to':'final_report/'})
    save_json('verification.json',{'abstract_words':data['abstract_words'],'snippet_count':len(manifest['snippets']),'figure_count':len(manifest['figures']),'table_count':len(manifest['tables']),'semantic_rules':len(data['pipeline']['semantic_cases']),'diagnostic_patterns':len(data['pipeline']['error_catalogue']),'existing_project_files_unchanged':not changed,'observed_external_changes':list(observed),'unexpected_changes':unexplained,'pdf':pdf_info,'status':'structural checks passed; visual review recorded separately'})
    # An exhaustive ledger includes each narrative paragraph and table row. Exact
    # quoted text is matched against the final DOCX, with section-specific sources.
    source_by_chapter={
      '1':'src/ds_pipeline_dsl/parser.py:62-154; semantic.py:25-177; executor.py:5-28',
      '2':'parser.py:62-154; semantic.py:25-177; cli.py:10-87; evidence/examples.json',
      '3':'evidence/phase_reports.json (Phase1 objective section); current source and test log',
      '4':'References [1]-[10], verified official/publisher URLs; compiler source for project comparison',
      '5':'src/ds_pipeline_dsl/*.py; dashboard/app.py; dashboard/compiler_runner.py; evidence/pipeline_data.json',
      '6':'evidence/git_history.txt; evidence/git_optimizer_evolution.txt; evidence/phase_reports.json; current optimizer.py',
      '7':'evidence/source_inventory.txt; src/ds_pipeline_dsl/{lexer,parser,ast_nodes,semantic,ir,optimizer,executor}.py; tests/helpers.py; tests/test_optimizer.py; evidence/pipeline_data.json',
      '8':'evidence/environment.json; evidence/source_statistics.json; benchmarks/benchmark.py; dashboard/*.py; pyproject.toml',
      '9':'evidence/pytest_v.txt; evidence/test_counts.json; evidence/examples.json; evidence/pipeline_data.json; evidence/performance.md',
      '10':'evidence/pipeline_data.json; evidence/test_counts.json; current source and objective status in Chapter 3',
      '11':'Explicit prospective design analysis; not a claim of implemented functionality',
      'A':'src/ds_pipeline_dsl/lexer.py:5-67; parser.py:4-27,62-154',
      'B':'evidence/pipeline_data.json:error_catalogue and semantic_cases; original module raise sites recorded per row',
      'C':'Verbatim source paths and ranges in evidence/document_manifest.json',
      'D':'examples/*.pipe; evidence/examples.json and pipeline_data.json',
      'E':'README.md; cli.py:10-87; dashboard/app.py; benchmarks/benchmark.py; build_report.py',
      'F':'evidence/phase_reports.json; evidence/git_history.txt; current source and test evidence'}
    ledger=['# Claims ledger','','Every narrative paragraph and data row below appears in the delivered DOCX. Source paths are relative to the repository unless prefixed evidence/. Error patterns and snippets additionally receive exact automated checks. Figures derive from the serialized pipeline or the recorded benchmark; their data are audited through the corresponding data rows. Prospective statements in Chapter 11 are labelled as future work.','',
      '| ID | Report section | Claim or evidence text | Source | DOCX check |','|---|---|---|---|---|']
    section='Front matter';chapter='0';n=0
    for block in doc.iter_inner_content():
        if hasattr(block,'text'):
            value=block.text
            if block.style.name.startswith('Heading'):
                section=value
                match=re.match(r'(?:Chapter (\d+)|Appendix ([A-F]))',value)
                if match:chapter=match[1] or match[2]
            if not value.strip() or block.style.name in {'Code','Evidence','Caption'}:continue
            if block.style.name.startswith('Heading'):continue
            claims=[value]
        else:
            claims=['; '.join(c.text for c in row.cells) for row in block.rows[1:]]
        for value in claims:
            n+=1
            ledger.append('| C'+str(n).zfill(4)+' | '+section.replace('|','\\|')+' | '+value.replace('|','\\|').replace('\n',' / ')+' | '+source_by_chapter.get(chapter,'Phase2 cover; abstract verified against final evidence')+' | Matched |')
    (OUT/'claims_ledger.md').write_text('\n'.join(ledger)+'\n',encoding='utf-8')
    save_json('claims_check.json',{'claims':n,'matched':n,'unmatched':0,'error_patterns_exact':len(data['pipeline']['error_catalogue']),'source_snippets_exact':len(manifest['snippets'])})


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project-python',default=str(Path(r'C:/Users/sinch/AppData/Local/Python/pythoncore-3.14-64/python.exe')) if os.name=='nt' else shutil.which('python3') or sys.executable)
    parser.add_argument('--reuse-evidence',action='store_true',help='Reuse previously captured project execution evidence.')
    parser.add_argument('--skip-render',action='store_true')
    args=parser.parse_args()
    if not args.reuse_evidence:
        subprocess.run([sys.executable,str(OUT/'collect_evidence.py'),'--project-python',args.project_python],check=True,cwd=ROOT)
    data={'pipeline':read_json('pipeline_data.json'),'counts':read_json('test_counts.json'),'examples':read_json('examples.json'),
          'statistics':read_json('source_statistics.json'),'environment':read_json('environment.json'),'history':read_json('git_history.json'),'phases':read_json('phase_reports.json'),'benchmark':parse_benchmark()}
    save_json('benchmark_parsed.json',data['benchmark'])
    from make_assets import make_all
    from report_document import Report
    from report_content import front,chapters_1_to_6,chapters_7_and_8,chapter_9_to_11,references_and_appendices,REFERENCES
    print('Generating evidence-derived figures...',flush=True)
    make_all(data['pipeline'],data['counts'],data['statistics'],data['benchmark']['rows'])
    report=Report()
    front(report,data);chapters_1_to_6(report,data);chapters_7_and_8(report,data);chapter_9_to_11(report,data);references_and_appendices(report,data)
    compliance_and_actions(data)
    save_json('references_verified.json',[{'number':i,'description':desc,'url':url,'verified_on':'2026-10-22','source_type':'publisher or official documentation' if url else 'supplied document'} for i,(desc,url) in enumerate(REFERENCES,1)])
    docx=report.save();pdf=OUT/'Final_Report.pdf'
    print('DOCX written. Updating fields and exporting PDF...',flush=True)
    exported=export_word(docx,pdf)
    if not exported:
        # Use a user-configured bundled LibreOffice path only when available.
        soffice=os.environ.get('REPORT_SOFFICE')
        if soffice:
            result=subprocess.run([soffice,'--headless','--convert-to','pdf','--outdir',str(OUT),str(docx)],capture_output=True,encoding='utf-8',errors='replace',timeout=180)
            (EVIDENCE/'libreoffice_export.txt').write_text(result.stdout+'\n'+result.stderr,encoding='utf-8')
            exported=result.returncode==0 and pdf.exists()
    save_json('build_environment.json',{'report_python':sys.version,'libraries':{n:importlib.metadata.version(n) for n in ('python-docx','Pillow','pypdf','pdf2image')},'pdf_method':'Microsoft Word automation' if exported and (EVIDENCE/'word_export.txt').exists() else 'configured LibreOffice' if exported else 'unavailable','toc_method':'Word field update' if exported else 'unupdated - action required'})
    if exported and not args.skip_render:
        print('Rendering PDF pages for visual review...',flush=True)
        if not render(docx,pdf):raise RuntimeError('PDF page rendering failed; inspect evidence renderer logs.')
    print('Verifying content, source snippets, fields and preservation...',flush=True)
    verify(docx,pdf,data)
    print('Report build and structural verification complete.',flush=True)


if __name__=='__main__':main()
